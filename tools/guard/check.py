#!/usr/bin/env python3
"""Locus guard: mechanical enforcement of the rules in AGENTS.md.

    python tools/guard/check.py --fast          # pre-commit (seconds)
    python tools/guard/check.py --full          # pre-push / CI / end of every task
    python tools/guard/check.py --full --deep   # also import-smoke + TypeScript typecheck
    python tools/guard/check.py --adopt-top-level   # one-time: accept today's top-level layout
    python tools/guard/check.py --update-spec-baseline   # human only: accept current spec defects

Exit code 1 if any FAIL. WARN never blocks unless --strict.
Nothing here calls an LLM. If this script says FAIL, the work is not done.
"""
from __future__ import annotations

import argparse
import ast
import fnmatch
import json
import os
import re
import stat
import subprocess
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import locus_common as C  # noqa: E402
from locus_common import Finding  # noqa: E402

CODE_EXTS = {".py", ".ts", ".tsx", ".js", ".mjs", ".cjs", ".sh", ".ps1", ".bat", ".cmd", ".service",
             ".yml", ".yaml", ".toml", ".json"}


class Ctx:
    def __init__(self, root: Path, cfg: dict, hook: bool, staged_only: bool, deep: bool):
        self.root, self.cfg, self.hook, self.staged_only, self.deep = root, cfg, hook, staged_only, deep
        w = cfg.get("walk", {})
        self.skip_dirs = set(w.get("skip_dirs", []))
        self.upstream_dirs = list(w.get("upstream_dirs", ["packages"]))
        self.scan_exclude = list(w.get("scan_exclude", []))
        self._code: list[Path] | None = None
        self._docs: list[Path] | None = None

    # -- file sets ----------------------------------------------------------
    def _scan(self, want_docs: bool) -> list[Path]:
        out = []
        skip_rel = self.upstream_dirs + self.scan_exclude
        for p in C.iter_files(self.root, self.skip_dirs, skip_rel):
            r = C.rel(self.root, p)
            if r.startswith("vault/") or r == "vault":
                continue
            if not C.is_text_file(p):
                continue
            is_doc = p.suffix.lower() == ".md"
            if want_docs == is_doc:
                if want_docs or p.suffix.lower() in CODE_EXTS or p.name.startswith(("Dockerfile", "docker-compose")):
                    out.append(p)
        return out

    def code_files(self) -> list[Path]:
        if self._code is None:
            self._code = self._scan(False)
        return self._code

    def doc_files(self) -> list[Path]:
        if self._docs is None:
            self._docs = self._scan(True)
        return self._docs


def glob_match(rel_path: str, name: str, globs: list[str]) -> bool:
    return any(fnmatch.fnmatch(name, g) or fnmatch.fnmatch(rel_path, g) for g in globs)


def line_of(text: str, pos: int) -> int:
    return text.count("\n", 0, pos) + 1


# ============================================================ individual checks
def check_forbidden_paths(ctx: Ctx) -> list[Finding]:
    out = []
    for p in ctx.cfg.get("layout", {}).get("forbidden", []):
        if (ctx.root / p).exists():
            out.append(Finding("FAIL", "forbidden-path", f"`{p}` exists but is forbidden by guard.toml. "
                               "Delete it in a dedicated `discard:` commit (git rm -r), do not recreate it.", p))
    return out


def reparse_kind(p: Path) -> str | None:
    try:
        st = os.lstat(p)
    except OSError:
        return None
    if stat.S_ISLNK(st.st_mode):
        return "symlink"
    tag = getattr(st, "st_reparse_tag", 0)
    mount = getattr(stat, "IO_REPARSE_TAG_MOUNT_POINT", 0xA0000003)
    sym = getattr(stat, "IO_REPARSE_TAG_SYMLINK", 0xA000000C)
    if tag in (mount, sym):  # OneDrive cloud placeholders use other tags and are NOT flagged
        return "junction" if tag == mount else "symlink"
    return None


def check_reparse(ctx: Ctx) -> list[Finding]:
    out = []
    allow = ctx.cfg.get("reparse", {}).get("allow", [])
    skip = set(C.DEFAULT_SKIP_DIRS) | ctx.skip_dirs
    for dirpath, dirnames, filenames in os.walk(ctx.root, followlinks=False):
        keep = []
        for d in dirnames:
            p = Path(dirpath) / d
            if d in skip:
                continue
            kind = reparse_kind(p)
            if kind:
                r = C.rel(ctx.root, p)
                if r not in allow:
                    out.append(Finding("FAIL", "reparse-point", f"{kind} at `{r}`. A junction/symlink makes two paths "
                                       "the same files (this caused the core/sidecar overwrite). Remove it.", r))
                continue  # never descend
            keep.append(d)
        dirnames[:] = keep
        for fn in filenames:
            p = Path(dirpath) / fn
            kind = reparse_kind(p)
            if kind:
                r = C.rel(ctx.root, p)
                if r not in allow:
                    out.append(Finding("FAIL", "reparse-point", f"{kind} file at `{r}`.", r))
    return out


def check_forbidden_text(ctx: Ctx) -> list[Finding]:
    out = []
    rules = []
    for r in ctx.cfg.get("forbidden_text", []):
        rules.append((r, re.compile(r["pattern"], re.M)))
    if not rules:
        return out
    for p in ctx.code_files():
        r_rel = C.rel(ctx.root, p)
        text = None
        for rule, rx in rules:
            if not glob_match(r_rel, p.name, rule["globs"]):
                continue
            if text is None:
                text = C.read_text(p)
                if text is None:
                    break
            hits = list(rx.finditer(text))
            for m in hits[:3]:
                out.append(Finding(rule.get("severity", "FAIL"), rule["id"], rule["message"],
                                   f"{r_rel}:{line_of(text, m.start())}"))
            if len(hits) > 3:
                out.append(Finding("INFO", rule["id"], f"... {len(hits) - 3} more in this file", r_rel))
    return out


def check_stale_paths(ctx: Ctx) -> list[Finding]:
    out = []
    rules = [(r, re.compile(r["pattern"])) for r in ctx.cfg.get("stale_paths", [])]
    if not rules:
        return out
    for p in ctx.code_files():
        text = C.read_text(p)
        if text is None:
            continue
        for rule, rx in rules:
            m = rx.search(text)
            if m:
                out.append(Finding("FAIL", rule["id"], f"Spec-only path `{m.group(0)}` used in code. "
                                   f"Canonical home: `{rule['canonical']}`.", f"{C.rel(ctx.root, p)}:{line_of(text, m.start())}"))
    for p in ctx.doc_files():
        text = C.read_text(p)
        if text is None:
            continue
        for rule, rx in rules:
            hits = list(rx.finditer(text))
            if hits:
                out.append(Finding("WARN", rule["id"], f"{len(hits)} reference(s) to a non-canonical path; "
                                   f"canonical is `{rule['canonical']}`.", f"{C.rel(ctx.root, p)}:{line_of(text, hits[0].start())}"))
    return out


BUILTIN_TOP = {".git", ".gitignore", ".gitattributes", ".github", ".githooks", "AGENTS.md", "CLAUDE.md",
               "tools", "tracker", "tests", "docs", "packages", "sidecar", "locus-extension", "locus_cli"}


def top_level_entries(root: Path) -> list[str]:
    return sorted(e for e in os.listdir(root) if e not in BUILTIN_TOP)


def check_layout(ctx: Ctx) -> list[Finding]:
    out = []
    lay = ctx.cfg.get("layout", {})
    for d in lay.get("required_dirs", []):
        if not (ctx.root / d).is_dir():
            out.append(Finding("FAIL", "layout-required", f"Required folder `{d}` is missing.", d))
    for f in lay.get("required_files", []):
        if not (ctx.root / f).is_file():
            out.append(Finding("FAIL", "layout-required", f"Required file `{f}` is missing.", f))
    allowed = set(lay.get("allowed_top_level", []))
    forb = set(lay.get("forbidden", []))
    extra = [e for e in top_level_entries(ctx.root) if e not in allowed and e not in forb]
    if not allowed:
        out.append(Finding("WARN", "layout-baseline", "Top-level baseline not adopted yet. Review the folder listing, then run: "
                           "python tools/guard/check.py --adopt-top-level   (unknown top-level entries: "
                           + ", ".join(extra[:25]) + (" ..." if len(extra) > 25 else "") + ")"))
    else:
        for e in extra:
            out.append(Finding("FAIL", "layout-top-level", f"New top-level entry `{e}` was not accepted. Agents may not create "
                               "top-level folders/files. Move it under the right canonical folder, or (human) adopt it.", e))
    return out


def adopt_top_level(root: Path, cfg_path: Path, cfg: dict) -> int:
    forb = set(cfg.get("layout", {}).get("forbidden", []))
    bad = [e for e in os.listdir(root) if e in forb]
    if bad:
        print(f"Refusing to adopt: forbidden entries exist: {bad}. Delete them first.")
        return 2
    entries = top_level_entries(root)
    body = "allowed_top_level = [" + ", ".join(json.dumps(e) for e in entries) + "]"
    text = cfg_path.read_text(encoding="utf-8")
    new, n = re.subn(r"allowed_top_level\s*=\s*\[[^\]]*\]", lambda _m: body, text, count=1, flags=re.S)
    if n != 1:
        print("Could not find `allowed_top_level = [...]` in guard.toml")
        return 2
    cfg_path.write_text(new, encoding="utf-8", newline="\n")
    print(f"Adopted {len(entries)} top-level entries into guard.toml:")
    for e in entries:
        print("   ", e)
    print("Review `git diff tools/guard/guard.toml`. Anything here that should not exist: remove it, then re-run.")
    return 0


VAULT_TOKEN = re.compile(r"(?<![\w.\-])(0\d-[A-Za-z][A-Za-z0-9\-]*)(?![\w])")


def vault_suggest(cfg: dict, tok: str) -> str:
    aliases = {k.lower(): v for k, v in cfg.get("vault", {}).get("aliases", {}).items()}
    return aliases.get(tok.lower(), "")


def check_vault(ctx: Ctx) -> list[Finding]:
    out = []
    v = ctx.cfg.get("vault", {})
    canon = {c.lower() for c in v.get("canonical", [])}
    if not canon:
        return out
    vroot = ctx.root / v.get("root", "vault")
    if vroot.is_dir():
        for d in sorted(vroot.iterdir()):
            if d.is_dir() and not d.name.startswith("."):
                if d.name.lower() not in canon:
                    s = vault_suggest(ctx.cfg, d.name)
                    hint = f" Move its contents to `{s}`." if s else " Not a canonical vault folder."
                    out.append(Finding("FAIL", "vault-layout", f"Vault folder `{d.name}` is not canonical.{hint}", f"{v.get('root','vault')}/{d.name}"))
    for p in ctx.code_files():
        if p.suffix.lower() not in {".py", ".ts", ".tsx", ".js", ".sh", ".ps1", ".bat", ".cmd", ".yml", ".yaml", ".json", ".toml"}:
            continue
        text = C.read_text(p)
        if text is None:
            continue
        seen = set()
        for m in VAULT_TOKEN.finditer(text):
            tok = m.group(1)
            if tok.lower() in canon or tok.lower() in seen:
                continue
            seen.add(tok.lower())
            s = vault_suggest(ctx.cfg, tok)
            hint = f" Use `{s}`." if s else " Not a canonical vault folder."
            out.append(Finding("FAIL", "vault-name", f"Non-canonical vault folder name `{tok}`.{hint}", f"{C.rel(ctx.root, p)}:{line_of(text, m.start())}"))
    return out


def check_upstream(ctx: Ctx) -> list[Finding]:
    up = ctx.cfg.get("upstream", {})
    lock_path = ctx.root / up.get("lock_file", "tools/guard/upstream.lock.json")
    if not (ctx.root / ".git").exists():
        return [Finding("WARN", "upstream-clean", "Not a git repo; cannot verify upstream Pi is untouched.")]
    if not lock_path.is_file():
        return [Finding("FAIL", "upstream-clean", "No upstream pin. Run: python tools/guard/pin_upstream.py "
                        "(records the last clean Pi commit so edits to packages/ can be detected).")]
    lock = json.loads(lock_path.read_text(encoding="utf-8"))
    base, paths = lock["baseline_commit"], lock.get("paths", ctx.upstream_dirs)
    rc, _ = C.git(ctx.root, "cat-file", "-e", f"{base}^{{commit}}")
    if rc != 0:
        return [Finding("FAIL", "upstream-clean", f"Pinned baseline commit {base[:10]} is not in this repo's history.")]
    rc, diff = C.git(ctx.root, "diff", "--name-only", base, "--", *paths)
    rc2, untracked = C.git(ctx.root, "ls-files", "--others", "--exclude-standard", "--", *paths)
    changed = [x for x in (diff.splitlines() + untracked.splitlines()) if x.strip()]
    allow = up.get("allow_modified", [])
    changed = [c for c in changed if not any(fnmatch.fnmatch(c, a) for a in allow)]
    out = []
    for c in changed[:50]:
        out.append(Finding("FAIL", "upstream-clean", "Upstream Pi file differs from the pinned baseline. Locus must be an "
                           "overlay (locus-extension/ + sidecar/); revert this file.", c))
    if len(changed) > 50:
        out.append(Finding("FAIL", "upstream-clean", f"... and {len(changed) - 50} more upstream files differ."))
    return out


def _defined_names(tree: ast.AST) -> tuple[dict[str, ast.AST], set[str]]:
    defined: dict[str, ast.AST] = {}
    imported: set[str] = set()
    for node in getattr(tree, "body", []):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            defined[node.name] = node
        elif isinstance(node, ast.Assign):
            for t in node.targets:
                if isinstance(t, ast.Name):
                    defined[t.id] = node
        elif isinstance(node, (ast.Import, ast.ImportFrom)):
            for a in node.names:
                imported.add((a.asname or a.name).split(".")[0])
    return defined, imported


def check_must_define(ctx: Ctx) -> list[Finding]:
    out = []
    for rel_path, names in ctx.cfg.get("must_define", {}).items():
        p = ctx.root / rel_path
        if not p.is_file():
            out.append(Finding("FAIL", "must-define", f"`{rel_path}` is missing.", rel_path))
            continue
        try:
            tree = ast.parse(p.read_text(encoding="utf-8", errors="replace"))
        except SyntaxError as e:
            out.append(Finding("FAIL", "must-define", f"`{rel_path}` has a syntax error: {e.msg}", f"{rel_path}:{e.lineno}"))
            continue
        defined, imported = _defined_names(tree)
        for n in names:
            if n in defined:
                node = defined[n]
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and len(node.body) < 3:
                    out.append(Finding("WARN", "must-define", f"`{n}` in `{rel_path}` has a very short body ({len(node.body)} statements). Stub?", rel_path))
            elif n in imported:
                out.append(Finding("FAIL", "must-define", f"`{n}` in `{rel_path}` is only imported/re-exported, not defined. "
                                   "The real implementation may have been overwritten.", rel_path))
            else:
                out.append(Finding("FAIL", "must-define", f"`{n}` is not defined in `{rel_path}`.", rel_path))
    return out


def check_pi_bridge(ctx: Ctx) -> list[Finding]:
    pb = ctx.cfg.get("pi_bridge")
    if not pb:
        return []
    p = ctx.root / pb["file"]
    if not p.is_file():
        return []  # reported by must-define / layout
    text = p.read_text(encoding="utf-8", errors="replace")
    out = []
    for tok in pb.get("required_tokens", []):
        if tok not in text:
            out.append(Finding("FAIL", "pi-bridge-flags", f"`{pb['file']}` does not pass `{tok}` to Pi.", pb["file"]))
    for group in pb.get("required_any_of", []):
        if not any(g in text for g in group):
            out.append(Finding("FAIL", "pi-bridge-flags", f"`{pb['file']}` must pass one of {group} to Pi, otherwise Pi loads "
                               "AGENTS.md/CLAUDE.md (developer rules) into every runtime turn.", pb["file"]))
    return out


def check_syntax(ctx: Ctx) -> list[Finding]:
    out = []
    for pkg in ctx.cfg.get("python", {}).get("packages", []):
        base = ctx.root / pkg
        if not base.exists():
            continue
        for p in C.iter_files(base, ctx.skip_dirs):
            if p.suffix != ".py":
                continue
            try:
                ast.parse(p.read_text(encoding="utf-8", errors="replace"))
            except SyntaxError as e:
                out.append(Finding("FAIL", "python-syntax", f"Syntax error: {e.msg}", f"{C.rel(ctx.root, p)}:{e.lineno}"))
    return out


def check_agent_files(ctx: Ctx) -> list[Finding]:
    out = []
    a = ctx.cfg.get("agents", {})
    agents = ctx.root / "AGENTS.md"
    if agents.is_file():
        text = agents.read_text(encoding="utf-8", errors="replace")
        size = len(text.encode("utf-8"))
        if a.get("agents_md_marker") and a["agents_md_marker"] not in text[:600]:
            out.append(Finding("FAIL", "agents-md", f"AGENTS.md does not carry `{a['agents_md_marker']}`. It is probably Pi's upstream "
                               "'Development Rules' file. Move it: git mv AGENTS.md docs/upstream/pi-AGENTS.md, then install the Locus AGENTS.md.", "AGENTS.md"))
        if size > a.get("agents_md_max_bytes", 9000):
            out.append(Finding("FAIL", "agents-md", f"AGENTS.md is {size} bytes (limit {a.get('agents_md_max_bytes')}). It is injected on every turn; "
                               "long rules get ignored. Move detail into docs/agent/.", "AGENTS.md"))
    claude = ctx.root / "CLAUDE.md"
    if claude.is_file():
        if "@AGENTS.md" not in claude.read_text(encoding="utf-8", errors="replace"):
            out.append(Finding("FAIL", "claude-md", "CLAUDE.md must import AGENTS.md (a line containing @AGENTS.md) so Claude Code and Antigravity share one rule set.", "CLAUDE.md"))
    elif agents.is_file():
        out.append(Finding("WARN", "claude-md", "CLAUDE.md missing (Claude Code will not see AGENTS.md).", "CLAUDE.md"))
    if a.get("forbid_root_gemini_md") and (ctx.root / "GEMINI.md").exists():
        out.append(Finding("FAIL", "gemini-md", "Root GEMINI.md exists next to AGENTS.md; Antigravity documents no precedence between them. "
                           "Keep one rule source: delete GEMINI.md (put Gemini-specific notes in AGENTS.md).", "GEMINI.md"))
    seg_dir = ctx.root / "docs" / "agent"
    if seg_dir.is_dir():
        for p in sorted(seg_dir.rglob("*.md")):
            size = p.stat().st_size
            if size > a.get("segment_max_bytes", 7500) and "09-spec-errata" not in p.name and "prompts" not in p.parts:
                out.append(Finding("WARN", "segment-size", f"{C.rel(ctx.root, p)} is {size} bytes (limit {a.get('segment_max_bytes')}). "
                                   "Segments must stay concentrated: split it.", C.rel(ctx.root, p)))
    return out


def tracked_candidates(ctx: Ctx) -> list[str]:
    if ctx.staged_only:
        rc, out = C.git(ctx.root, "diff", "--cached", "--name-only", "--diff-filter=ACM")
    else:
        rc, out = C.git(ctx.root, "ls-files", "-co", "--exclude-standard")
    if rc != 0:
        return []
    return [x for x in out.splitlines() if x.strip()]


def check_secrets(ctx: Ctx) -> list[Finding]:
    s = ctx.cfg.get("secrets", {})
    pats = [(x["id"], re.compile(x["pattern"])) for x in s.get("patterns", [])]
    allow = s.get("allow_files", [])
    out = []
    for r in tracked_candidates(ctx):
        name = os.path.basename(r)
        if glob_match(r, name, allow) or r.startswith(("vault/", "packages/")):
            continue
        p = ctx.root / r
        if not p.is_file() or p.suffix.lower() in {".xlsx", ".png", ".jpg", ".zip", ".db", ".sqlite", ".lock"}:
            continue
        if name == ".env" or (name.startswith(".env.") and name != ".env.example"):
            out.append(Finding("FAIL", "secrets", f"`{r}` is a real env file inside the repo. Secrets live in ~/.locus/.env, never in the repo.", r))
            continue
        text = C.read_text(p, 1_000_000)
        if not text:
            continue
        for sid, rx in pats:
            m = rx.search(text)
            if m and "guard:allow-secret" not in text[max(0, m.start() - 200): m.end() + 200]:
                out.append(Finding("FAIL", "secrets", f"Looks like a {sid} credential. Rotate it NOW if it was ever pushed or pasted anywhere, then remove it.", f"{r}:{line_of(text, m.start())}"))
    return out


def check_gitignore(ctx: Ctx) -> list[Finding]:
    gi = ctx.root / ".gitignore"
    req = ctx.cfg.get("gitignore", {}).get("required_lines", [])
    if not req:
        return []
    have = set()
    if gi.is_file():
        have = {ln.strip() for ln in gi.read_text(encoding="utf-8", errors="replace").splitlines() if ln.strip() and not ln.startswith("#")}
    missing = [r for r in req if r not in have]
    if missing:
        return [Finding("WARN", "gitignore", "Missing .gitignore lines: " + ", ".join(missing), ".gitignore")]
    return []


SECTION_TOKENS = re.compile(r"^(\d+\.\d+[a-z]?)$")


def spec_issues(ctx: Ctx) -> dict[str, str]:
    """Deterministic defects in the spec file itself. key -> human message."""
    issues: dict[str, str] = {}
    sp_path = ctx.root / ctx.cfg["repo"]["spec_path"]
    if not sp_path.is_file():
        return {"spec-missing": f"Spec file not found at {ctx.cfg['repo']['spec_path']}"}
    spec = C.parse_spec(sp_path)
    heads, text, lines = spec["headings"], spec["text"], spec["lines"]
    # 1. duplicated exact section ids
    seen: dict[str, list[dict]] = {}
    for h in heads:
        for tok in re.split(r"\s*[+,]\s*", h["idpart"].replace(C.SECT, "")):
            tok = tok.strip()
            if SECTION_TOKENS.match(tok):
                seen.setdefault(tok, []).append(h)
    for tok, hs in seen.items():
        if len(hs) > 1:
            issues[f"dup-id:{tok}"] = f"Section id {C.SECT}{tok} is used by {len(hs)} different sections: " + "; ".join(f"line {h['line']} {h['title'][:40]}" for h in hs)
    # 2. conflicting TWS definitions
    tws = [(i + 1, ln.strip()) for i, ln in enumerate(lines) if re.search(r"\bTWS\s*=\s*\(", ln)]
    if len({re.sub(r"\s+", "", t[1]) for t in tws}) > 1:
        issues["tws-formula-conflict"] = "TWS has conflicting formulas at lines " + ", ".join(str(t[0]) for t in tws)
    # 3. vault names / stale paths in the spec
    canon = {c.lower() for c in ctx.cfg.get("vault", {}).get("canonical", [])}
    for tok in sorted({m.group(1) for m in VAULT_TOKEN.finditer(text)}):
        if tok.lower() not in canon:
            issues[f"vault-name:{tok}"] = f"Spec uses non-canonical vault folder `{tok}`"
    for r in ctx.cfg.get("stale_paths", []):
        n = len(re.findall(r["pattern"], text))
        if n:
            issues[f"stale-path:{r['id']}"] = f"Spec uses non-canonical path ({r['id']}) {n} time(s); canonical `{r['canonical']}`"
    # 4. missing principles
    for pr in ctx.cfg.get("spec_lint", {}).get("required_principles", []):
        if not re.search(pr["pattern"], text, re.I):
            issues[f"principle-missing:{pr['id']}"] = f"Spec never states: {pr['why']}"
    # 5. self-reported feature count vs real headings
    m = re.search(r"(\d+)\s+features?\s+total", text)
    if m and abs(int(m.group(1)) - len(heads)) > max(3, len(heads) // 10):
        issues[f"feature-count-claim:{m.group(1)}"] = f"Spec claims {m.group(1)} features total but has {len(heads)} feature sections"
    # 6. index tables that give the same id different names
    first_feature = heads[0]["line"] if heads else len(lines)
    names: dict[str, set[str]] = {}
    for ln in lines[:first_feature]:
        mm = re.match(r"^\|\s*(\d+\.\d+)\s*\|\s*([^|]+?)\s*\|", ln)
        if mm:
            names.setdefault(mm.group(1), set()).add(C.norm_key(mm.group(2)))
    def _tok(x: str) -> set[str]:
        return set(re.findall(r"[a-z0-9]+", x))

    for k, v in names.items():
        vs = sorted(v)
        clash = any(len(_tok(a) & _tok(b)) / max(1, len(_tok(a) | _tok(b))) < 0.5 for i, a in enumerate(vs) for b in vs[i + 1:])
        if len(vs) > 1 and clash:
            issues[f"table-conflict:{k}"] = f"Index tables name {C.SECT}{k} differently: " + " / ".join(vs)
    return issues


def check_spec_lint(ctx: Ctx) -> list[Finding]:
    cfg = ctx.cfg.get("spec_lint", {})
    base_path = ctx.root / cfg.get("baseline_file", "tools/guard/spec_lint.baseline.json")
    issues = spec_issues(ctx)
    if "spec-missing" in issues:
        return [Finding("WARN", "spec-lint", issues["spec-missing"])]
    baseline = set(json.loads(base_path.read_text(encoding="utf-8"))["known"]) if base_path.is_file() else set()
    out = []
    new = [k for k in issues if k not in baseline]
    fixed = [k for k in baseline if k not in issues]
    for k in new:
        out.append(Finding("FAIL", "spec-lint", f"NEW spec defect: {issues[k]}", k))
    for k in fixed:
        out.append(Finding("INFO", "spec-lint", f"Spec defect resolved; remove from baseline: {k}", k))
    known = [k for k in issues if k in baseline]
    if known:
        out.append(Finding("WARN", "spec-lint", f"{len(known)} KNOWN spec defects still open (docs/agent/09-spec-errata.md). "
                           "Agents must stop and ask when a task touches one."))
    return out


def update_spec_baseline(ctx: Ctx) -> int:
    issues = spec_issues(ctx)
    cfg = ctx.cfg.get("spec_lint", {})
    path = ctx.root / cfg.get("baseline_file", "tools/guard/spec_lint.baseline.json")
    C.json_dump({"note": "Accepted spec defects. HUMAN edits only. Shrink this list as the spec is fixed.",
                 "known": sorted(issues)}, path)
    print(f"Wrote {len(issues)} accepted spec defects to {C.rel(ctx.root, path)}")
    return 0


def check_tracker_integrity(ctx: Ctx) -> list[Finding]:
    fp = ctx.root / "tracker" / "features.yaml"
    if not fp.is_file():
        return [Finding("WARN", "tracker", "tracker/features.yaml not found (guardrail bundle not fully installed).")]
    out = []
    try:
        data = C.load_yaml(fp)
    except Exception as e:  # noqa: BLE001
        return [Finding("FAIL", "tracker", f"features.yaml does not parse: {e}")]
    feats = data.get("features", []) if isinstance(data, dict) else []
    ids, forbidden_keys = set(), {"status", "pct", "percent", "completion", "completion_pct", "level", "complete", "done", "progress"}
    for f in feats:
        fid = f.get("id", "?")
        if fid in ids:
            out.append(Finding("FAIL", "tracker", f"Duplicate feature id {fid}", fid))
        ids.add(fid)
        bad = forbidden_keys & set(f)
        if bad:
            out.append(Finding("FAIL", "tracker", f"{fid} sets {sorted(bad)}. Status is computed from evidence, never written. Delete those keys.", fid))
        for pth in (f.get("evidence", {}) or {}).get("files", []) or []:
            if ".." in Path(pth).parts or Path(pth).is_absolute():
                out.append(Finding("FAIL", "tracker", f"{fid}: evidence path escapes the repo: {pth}", fid))
    sp_path = ctx.root / ctx.cfg["repo"]["spec_path"]
    if sp_path.is_file():
        heads = C.parse_spec(sp_path)["headings"]
        keys = {h["key"] for h in heads}
        refs = set()
        for f in feats:
            for r in f.get("spec_refs", []) or []:
                k = C.norm_key(r)
                refs.add(k)
                if k not in keys:
                    out.append(Finding("FAIL", "tracker", f"{f.get('id')}: spec ref not found in the spec (renamed or deleted?): {r}", f.get("id", "")))
        excl_path = ctx.root / "tracker" / "spec_exclusions.yaml"
        excl = set()
        if excl_path.is_file():
            ex = C.load_yaml(excl_path) or {}
            excl = {C.norm_key(x["ref"]) for x in ex.get("exclusions", [])}
        for h in heads:
            if h["key"] not in refs and h["key"] not in excl:
                out.append(Finding("FAIL", "tracker", f"Spec section has no tracker feature: {C.display_heading(h)} (line {h['line']}). "
                                   "Add a feature for it, or list it in tracker/spec_exclusions.yaml with a reason.", f"spec:{h['line']}"))
    xlsx = ctx.root / ctx.cfg["repo"]["tracker_xlsx"]
    if xlsx.is_file():
        try:
            import openpyxl  # noqa: PLC0415
        except ImportError:
            return out + [Finding("WARN", "tracker-xlsx", "Cannot verify workbook integrity: pip install openpyxl")]
        try:
            wb = openpyxl.load_workbook(xlsx, data_only=False)
            if "_meta" not in wb.sheetnames:
                out.append(Finding("WARN", "tracker-xlsx", "Workbook is the old hand-maintained one (no _meta sheet). Replace it: python tools/tracker/export_xlsx.py", C.rel(ctx.root, xlsx)))
            else:
                meta = {r[0].value: r[1].value for r in wb["_meta"].iter_rows(min_row=1, max_col=2) if r[0].value}
                if C.grid_hash(wb["Feature Tracker"]) != meta.get("grid_sha"):
                    out.append(Finding("FAIL", "tracker-xlsx", "Workbook cells differ from what export_xlsx.py wrote: it was hand-edited (by you or an agent). "
                                       "Edit tracker/features.yaml instead, then regenerate.", C.rel(ctx.root, xlsx)))
        except Exception as e:  # noqa: BLE001
            out.append(Finding("WARN", "tracker-xlsx", f"Could not read workbook: {e}", C.rel(ctx.root, xlsx)))
    return out


def check_human_review(ctx: Ctx) -> list[Finding]:
    globs = ctx.cfg.get("human_review", {}).get("globs", [])
    if not globs or not (ctx.root / ".git").exists():
        return []
    if ctx.staged_only:
        rc, out = C.git(ctx.root, "diff", "--cached", "--name-only")
    else:
        rc, out = C.git(ctx.root, "diff", "--name-only", "HEAD")
        rc2, out2 = C.git(ctx.root, "ls-files", "--others", "--exclude-standard")
        out = out + "\n" + out2
    changed = [x for x in out.splitlines() if x.strip() and "__pycache__" not in x and not x.endswith(".pyc")
               and glob_match(x, os.path.basename(x), globs)]
    if not changed:
        return []
    allowed = os.environ.get("LOCUS_ALLOW_GUARD_EDIT") == "1"
    sev = "WARN" if (allowed or not ctx.hook) else "FAIL"
    msg = ("HUMAN REVIEW REQUIRED. Rule/guard/tracker files changed: " + ", ".join(sorted(set(changed))[:8]) +
           (" ..." if len(set(changed)) > 8 else "") +
           ("" if allowed else ". Agents must not edit these. If YOU made this change on purpose: set LOCUS_ALLOW_GUARD_EDIT=1 for this commit."))
    return [Finding(sev, "human-review", msg)]


def check_git_hygiene(ctx: Ctx) -> list[Finding]:
    if not (ctx.root / ".git").exists():
        return [Finding("FAIL", "git", "Not a git repository. Everything below depends on git history.")]
    out = []
    rc, br = C.git(ctx.root, "rev-parse", "--abbrev-ref", "HEAD")
    br = br.strip()
    if br in {"main", "master"} and C.git_dirty(ctx.root):
        out.append(Finding("WARN", "git", f"Uncommitted changes while on `{br}`. Agents work on `task/US-nn-<slug>` branches."))
    rc, remotes = C.git(ctx.root, "remote")
    if not remotes.strip():
        out.append(Finding("WARN", "git", "No git remote configured: the repo exists on one disk only. Add a private remote and push."))
    hp = C.git(ctx.root, "config", "core.hooksPath")[1].strip()
    if hp != ".githooks":
        out.append(Finding("WARN", "git", "Hooks not active. Run: git config core.hooksPath .githooks"))
    return out


def check_deep(ctx: Ctx) -> list[Finding]:
    out = []
    for mod in ctx.cfg.get("python", {}).get("smoke_imports", []):
        try:
            r = subprocess.run([sys.executable, "-c", f"import {mod}"], cwd=ctx.root, capture_output=True, text=True, timeout=90)
        except subprocess.TimeoutExpired:
            out.append(Finding("FAIL", "import-smoke", f"import {mod} timed out (90s)"))
            continue
        if r.returncode != 0:
            tail = (r.stderr or "").strip().splitlines()[-1:] or ["?"]
            out.append(Finding("FAIL", "import-smoke", f"import {mod} failed: {tail[0][:160]}"))
    tsconfig = ctx.root / "locus-extension" / "tsconfig.json"
    if tsconfig.is_file():
        try:
            r = subprocess.run("npx --no-install tsc --noEmit -p locus-extension", cwd=ctx.root, shell=True, capture_output=True, text=True, timeout=300)
            if r.returncode != 0:
                first = (r.stdout or r.stderr).strip().splitlines()[:3]
                out.append(Finding("FAIL", "ts-typecheck", "TypeScript errors in locus-extension: " + " | ".join(first)[:300]))
        except Exception as e:  # noqa: BLE001
            out.append(Finding("WARN", "ts-typecheck", f"Could not run tsc: {e}"))
    return out


# ================================================================== orchestration
FAST = [check_forbidden_paths, check_reparse, check_forbidden_text, check_stale_paths, check_layout, check_vault,
        check_must_define, check_pi_bridge, check_syntax, check_agent_files, check_secrets, check_human_review]
FULL_EXTRA = [check_upstream, check_spec_lint, check_tracker_integrity, check_gitignore, check_git_hygiene]


def run(ctx: Ctx, mode: str, only: set[str] | None) -> list[Finding]:
    checks = list(FAST) + (FULL_EXTRA if mode == "full" else [])
    if ctx.deep:
        checks.append(check_deep)
    findings: list[Finding] = []
    for fn in checks:
        name = fn.__name__.replace("check_", "").replace("_", "-")
        if only and name not in only:
            continue
        try:
            findings += fn(ctx)
        except Exception as e:  # noqa: BLE001
            findings.append(Finding("FAIL", name, f"guard crashed while running this check: {type(e).__name__}: {e}"))
    return findings


def main() -> int:
    C.utf8_stdio()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--fast", action="store_true", help="pre-commit subset (default)")
    g.add_argument("--full", action="store_true", help="everything except --deep")
    ap.add_argument("--deep", action="store_true", help="import smoke test + TypeScript typecheck")
    ap.add_argument("--hook", action="store_true", help="called from a git hook (human-review becomes blocking)")
    ap.add_argument("--staged", action="store_true", help="scan only staged files for secrets")
    ap.add_argument("--strict", action="store_true", help="WARN also fails")
    ap.add_argument("--all", action="store_true", help="print every finding (default: 8 per check)")
    ap.add_argument("--only", help="comma-separated check names, e.g. layout,vault")
    ap.add_argument("--root", help="repo root (default: inferred from this file)")
    ap.add_argument("--config", help="path to guard.toml")
    ap.add_argument("--json", help="write machine-readable report here")
    ap.add_argument("--adopt-top-level", action="store_true")
    ap.add_argument("--update-spec-baseline", action="store_true")
    args = ap.parse_args()

    root = Path(args.root).resolve() if args.root else Path(__file__).resolve().parents[2]
    cfg_path = Path(args.config) if args.config else root / "tools" / "guard" / "guard.toml"
    if not cfg_path.is_file():
        print(f"guard.toml not found at {cfg_path}")
        return 2
    cfg = C.load_toml(cfg_path)
    ctx = Ctx(root, cfg, args.hook, args.staged or args.hook, args.deep)

    if args.adopt_top_level:
        return adopt_top_level(root, cfg_path, cfg)
    if args.update_spec_baseline:
        return update_spec_baseline(ctx)

    mode = "full" if args.full else "fast"
    t0 = time.time()
    only = {x.strip() for x in args.only.split(",") if x.strip()} if args.only else None
    if only:
        names = lambda fns: {fn.__name__.replace("check_", "").replace("_", "-") for fn in fns}  # noqa: E731
        unknown = only - names(FAST + FULL_EXTRA + [check_deep])
        if unknown:
            print(f"unknown check name(s): {', '.join(sorted(unknown))}. Known: {', '.join(sorted(names(FAST + FULL_EXTRA)))}")
            return 2
        if only & names(FULL_EXTRA):
            mode = "full"  # a check that only runs in full mode must not be skipped silently
    findings = run(ctx, mode, only)

    order = {"FAIL": 0, "WARN": 1, "INFO": 2}
    findings.sort(key=lambda f: (order[f.severity], f.check))
    counts = {k: sum(1 for f in findings if f.severity == k) for k in order}
    print(f"LOCUS GUARD  mode={mode}{'+deep' if args.deep else ''}  root={root}")
    shown: dict[tuple[str, str], int] = {}
    for f in findings:
        key = (f.severity, f.check)
        shown[key] = shown.get(key, 0) + 1
        if args.all or shown[key] <= 8:
            print(f)
        elif shown[key] == 9:
            print(f"[{f.severity}] {f.check}: ... more (use --all)")
    verdict = "FAIL" if counts["FAIL"] or (args.strict and counts["WARN"]) else "PASS"
    print(f"\n{verdict}: {counts['FAIL']} FAIL, {counts['WARN']} WARN, {counts['INFO']} INFO  ({time.time() - t0:.1f}s)")
    if verdict == "FAIL":
        print("Not done. Fix every FAIL before you report completion.")
    report = {"verdict": verdict, "mode": mode, "only": sorted(only) if only else None, "counts": counts, "commit": C.git_head(root),
              "ran_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"), "findings": [f.as_dict() for f in findings]}
    out_path = Path(args.json) if args.json else root / "tracker" / ".cache" / "guard_report.json"
    try:
        C.json_dump(report, out_path)
    except OSError:
        pass
    return 1 if verdict == "FAIL" else 0


if __name__ == "__main__":
    sys.exit(main())
