#!/usr/bin/env python3
"""Convert the hand-maintained workbook into tracker/features.yaml (the source of truth).

    python tools/tracker/import_legacy.py                      # refuses to overwrite an existing features.yaml
    python tools/tracker/import_legacy.py --force              # Day-0 only: re-import from the workbook
    python tools/tracker/import_legacy.py --xlsx path/to.xlsx

What it does, deterministically (no LLM):
  * every legacy row becomes a feature (id F-0001...), its numbers are kept as `claimed` and NEVER used for scoring;
  * every claimed Key File becomes an evidence file to be tested for existence;
  * every spec section that no legacy row covers becomes a new feature (origin: spec-gap);
  * rows whose title does not match the spec section they point at are flagged mapping: low-confidence.
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import locus_common as C  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.stderr.write("openpyxl missing. Run: python -m pip install pyyaml openpyxl pytest\n")
    raise

STOP = {"the", "and", "of", "a", "for", "feature", "system", "engine", "layer", "to", "in", "on", "with", "or", "an", "via"}
SYN = {"tui": "terminal", "mcp": "model context protocol", "stt": "voice", "tts": "voice", "telegram": "channel",
       "cli": "commands", "doctor": "commands", "inherited": "", "docker": "docker"}


def toks(s: str) -> set[str]:
    out = set()
    for w in re.findall(r"[a-z0-9]+", s.lower()):
        w = SYN.get(w, w)
        for part in w.split():
            part = part.rstrip("s") if len(part) > 3 else part
            if part and part not in STOP:
                out.add(part)
    return out


def sim(a: str, b: str) -> float:
    ta, tb = toks(a), toks(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)


def norm_verdict(s: str | None) -> str:
    u = (s or "").upper()
    if "PRE-BUILT" in u or "PREBUILT" in u:
        return "KEEP-PREBUILT"
    for v in ("KEEP-PI", "MULTI-PROVIDER", "LOCUS-NEW", "AUGMENT"):
        if v in u:
            return v
    return "UNKNOWN"


def parse_key_files(raw: str | None) -> list[str]:
    if not raw:
        return []
    text = re.sub(r"\([^)]*\)", "", str(raw))
    out: list[str] = []
    for tok in re.split(r"[,;\n]", text):
        tok = tok.strip().strip("`").replace("\\", "/")
        if not tok or " " in tok:
            continue
        if "/" in tok or re.search(r"\.[A-Za-z0-9]{1,5}$", tok) or tok.startswith("."):
            if tok not in out:
                out.append(tok)
    return out


def layer_of(files: list[str]) -> str:
    kinds = set()
    for f in files:
        if f.startswith("packages/"):
            kinds.add("pi")
        elif f.startswith("locus-extension/"):
            kinds.add("extension")
        elif f.startswith(("sidecar/", "locus_cli/", "gateway/")):
            kinds.add("sidecar")
        else:
            kinds.add("infra")
    if not kinds:
        return "unknown"
    return kinds.pop() if len(kinds) == 1 else "mixed"


# Seeds: checks derived from specific, falsifiable claims in the two session reports.
TOOLS13 = ["locus_checkin", "locus_life", "locus_goals", "locus_memory", "locus_vault", "locus_brain", "locus_research",
           "locus_delegate_code", "locus_doctor", "locus_pageindex", "locus_diagram", "locus_capture", "locus_schedule"]
SEEDS = {
    "§1.3": {"symbols": [{"file": "locus-extension/src/tools/index.ts", "contains": TOOLS13}]},
    "ARCH-1": {"files": ["install.ps1", "install.sh", "install.bat", "locus.bat", "locus.cmd"]},
    "ARCH-3": {"files": ["sidecar/__init__.py"],
               "symbols": [{"file": "sidecar/__init__.py", "contains": ["__version__"]}],
               "guard": ["forbidden-path", "reparse-point", "py-import-core", "core-alias-shim", "core-module-run", "core-path-in-scripts"]},
}
EXTRA_ROWS = [{
    "title": "ARCH-6 Pi bridge (sidecar -> Pi CLI, used by cron + Telegram + conductor)",
    "category": "Infrastructure", "verdict": "LOCUS-NEW", "layer": "sidecar",
    "files": ["sidecar/pi_bridge.py", "sidecar/cron_runner.py"],
    "symbols": [{"file": "sidecar/pi_bridge.py", "name": "run_pi_turn"}, {"file": "sidecar/pi_bridge.py", "name": "get_node_executable"}],
    "guard": ["must-define", "pi-bridge-flags"],
    "origin": "session-report (Pi Agent Integration, Phase B)",
}]


def main() -> int:
    C.utf8_stdio()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx")
    ap.add_argument("--spec")
    ap.add_argument("--root")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    root = Path(args.root).resolve() if args.root else Path(__file__).resolve().parents[2]
    cfg = C.load_toml(root / "tools" / "guard" / "guard.toml")
    xlsx = Path(args.xlsx) if args.xlsx else root / cfg["repo"]["tracker_xlsx"]
    spec_path = Path(args.spec) if args.spec else root / cfg["repo"]["spec_path"]
    out_path = root / "tracker" / "features.yaml"
    if out_path.exists() and not args.force:
        print(f"{out_path} already exists. Refusing to overwrite (evidence you added would be lost). Use --force on Day 0 only.")
        return 2

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["Feature Tracker"]
    hdr = [c.value for c in ws[1]]
    legacy = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is None and row[1] is None:
            continue
        d = dict(zip(hdr, row))
        d["_row"] = i
        legacy.append(d)

    spec = C.parse_spec(spec_path)
    heads = spec["headings"]
    claimed_refs: dict[int, list[dict]] = {}
    feats: list[dict] = []
    used_heads: set[int] = set()

    for n, r in enumerate(legacy, start=1):
        sid = str(r["§ID"])
        ids = C.expand_ids(sid) if not sid.startswith("ARCH") else []
        is_range = bool(re.search(r"\d[-–]\d", sid.replace("§", "")))
        cands = [h for h in heads if set(h["ids"]) & set(ids)]
        name = str(r["Feature Name"])
        mapping, chosen = "none", []
        if cands:
            if is_range:
                chosen, mapping = cands, "grouped"
            else:
                best = max(cands, key=lambda h: sim(name, h["title"]))
                chosen, mapping = [best], "exact"
            best_sim = max(sim(name, h["title"]) for h in chosen)
            if best_sim < 0.15:
                mapping = "low-confidence"
        for h in chosen:
            used_heads.add(h["line"])
        files = parse_key_files(r.get("Key Files"))
        verdict = norm_verdict(r.get("Verdict"))
        inherited = verdict == "KEEP-PI" and bool(files) and all(f.startswith("packages/") for f in files)
        targets = sorted({t for h in chosen for t in h["targets"]})[:10]
        feat = {
            "id": f"F-{n:04d}",
            "title": name,
            "spec_refs": [C.display_heading(h) for h in chosen],
            "mapping": mapping,
            "origin": "legacy-tracker" if chosen or not sid.startswith("ARCH") else "user-requirement (not a spec section)",
            "category": r.get("Category"),
            "verdict": verdict,
            "source": r.get("Source Codebase / Stealing From"),
            "provider": None if str(r.get("Provider Interface")) in {"—", "None", "-"} else r.get("Provider Interface"),
            "layer": layer_of(files),
            "inherited": inherited,
            "manual": False,
            "claimed": {
                "legacy_id": sid, "legacy_row": r["_row"], "pi_status": r.get("Pi Agent Status"), "sidecar_status": r.get("Sidecar Status"),
                "pct": r.get("Completion %"), "remaining": r.get("Remaining Work"), "notes": r.get("Notes / Issues"),
                "key_files_raw": r.get("Key Files"),
            },
            "spec_targets": targets,
            "spec_sha": {C.display_heading(h): h["sha"] for h in chosen},
            "evidence": {"files": files, "symbols": [], "guard": [], "unit": [], "integration": []},
            "depends_on": [],
        }
        seed = SEEDS.get(sid)
        if seed:
            ev = feat["evidence"]
            for k, v in seed.items():
                for item in v:
                    if item not in ev.setdefault(k, []):
                        ev[k].append(item)
        feats.append(feat)

    # extra rows derived from the session reports
    for ex in EXTRA_ROWS:
        n = len(feats) + 1
        feats.append({
            "id": f"F-{n:04d}", "title": ex["title"], "spec_refs": [], "mapping": "none", "origin": ex["origin"],
            "category": ex["category"], "verdict": ex["verdict"], "source": "Locus-original", "provider": None,
            "layer": ex["layer"], "inherited": False, "manual": False, "claimed": None, "spec_targets": [], "spec_sha": {},
            "evidence": {"files": ex["files"], "symbols": ex["symbols"], "guard": ex["guard"], "unit": [], "integration": []},
            "depends_on": [],
        })

    # spec sections that no legacy row covers
    gaps = [h for h in heads if h["line"] not in used_heads]
    for h in gaps:
        n = len(feats) + 1
        feats.append({
            "id": f"F-{n:04d}", "title": h["title"], "spec_refs": [C.display_heading(h)], "mapping": "exact",
            "origin": "spec-gap (the old tracker had no row for this section)",
            "category": None, "verdict": norm_verdict(h["verdict"]), "source": None, "provider": None, "layer": "unknown",
            "inherited": False, "manual": False, "claimed": None, "spec_targets": h["targets"][:10],
            "spec_sha": {C.display_heading(h): h["sha"]},
            "evidence": {"files": [], "symbols": [], "guard": [], "unit": [], "integration": []}, "depends_on": [],
        })

    header = (
        "# tracker/features.yaml -- SOURCE OF TRUTH for WHAT must exist and HOW to prove it.\n"
        "#\n"
        "#   * Status is COMPUTED by tools/tracker/verify.py from `evidence`. It is never written here.\n"
        "#     Keys like status / pct / complete are rejected by the guard.\n"
        "#   * `claimed` is what the old LLM-maintained workbook said. It is kept for comparison only.\n"
        "#   * To advance a feature, add evidence: files -> symbols -> unit (L3) -> integration (L4) -> proof + signoff (L5).\n"
        "#   * Do not read this whole file (thousands of lines). Use: python tools/tracker/show.py F-0007\n"
        "#   * Agents may ADD evidence for the feature they are working on. They may not delete or weaken existing evidence.\n"
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    body = C.yaml.safe_dump({"schema": 1, "spec_file": cfg["repo"]["spec_path"], "features": feats},
                            sort_keys=False, allow_unicode=True, width=110)
    out_path.write_text(header + body, encoding="utf-8", newline="\n")

    legacy_dir = root / "tracker" / "legacy"
    legacy_dir.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(xlsx, legacy_dir / "Locus_Implementation_Tracker.legacy.xlsx")
    C.json_dump({"imported_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"), "xlsx_sha256": C.sha256_file(xlsx), "rows": legacy},
                legacy_dir / "legacy_claims.json")
    for name, text in (("spec_exclusions.yaml", "# Spec sections deliberately NOT tracked. Human-owned.\n# exclusions:\n#   - ref: \"§1.8: Stream Diagnostics\"\n#     reason: \"...\"\nexclusions: []\n"),
                       ("overrides.yaml", "# Human-only caps: lower a feature's computed level with a reason.\n# F-0007: {cap_level: 2, by: Shivam, date: 2026-10-05, reason: \"...\"}\n{}\n"),
                       ("signoff.yaml", "# Human-only. A feature reaches L5 PROVEN only with a current proof file AND an entry here that names that proof.\n# record_proof.py prints the exact line to paste, e.g.:\n# F-0007: {by: Shivam, date: 2026-10-05, proof: \"3f9a1c2b7d10\", note: \"watched the cron run end to end\"}\n{}\n")):
        p = root / "tracker" / name
        if not p.exists():
            p.write_text(text, encoding="utf-8", newline="\n")

    by_map: dict[str, int] = {}
    for f in feats:
        by_map[f["mapping"]] = by_map.get(f["mapping"], 0) + 1
    print(f"legacy rows: {len(legacy)} | spec sections: {len(heads)} | gap features added: {len(gaps)} | total features: {len(feats)}")
    print("mapping:", by_map)
    print(f"wrote {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
