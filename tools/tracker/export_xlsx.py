#!/usr/bin/env python3
"""Regenerate Locus_Implementation_Tracker.xlsx from tracker/features.yaml + tracker/status.json.

    python tools/tracker/verify.py && python tools/tracker/export_xlsx.py

The workbook is a VIEW. Do not edit it: the guard compares its cells with what this script wrote
and fails if they differ (hand edits, by you or by an agent, are detected). Edit features.yaml.
The old hand-maintained workbook is backed up to tracker/legacy/ the first time this runs over it.
"""
from __future__ import annotations

import argparse
import shutil
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import locus_common as C  # noqa: E402

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.stderr.write("openpyxl missing. Run: python -m pip install pyyaml openpyxl pytest\n")
    raise

HEADERS = ["ID", "Spec section(s)", "Feature", "Category", "Verdict", "Source Codebase / Stealing From", "Provider Interface",
           "Layer", "Verified Level", "Verified %", "Claimed % (old tracker, unverified)", "Claim minus Evidence",
           "Remaining Work (derived)", "Flags", "Evidence Files", "Why Not Higher"]
assert len(HEADERS) == C.GRID_COLS
FILLS = {"NO_EVIDENCE": "F8CBAD", "FILES_PRESENT": "FFE699", "IMPLEMENTED": "FFF2CC", "UNIT_VERIFIED": "E2EFDA",
         "INTEGRATED": "C6E0B4", "PROVEN": "A9D08E", "INHERITED_OK": "BDD7EE"}
WIDTHS = [9, 34, 38, 14, 14, 26, 18, 10, 16, 10, 16, 12, 56, 28, 44, 56]


def num(v):
    if v is None:
        return None
    return int(v) if float(v).is_integer() else round(float(v), 1)


def main() -> int:
    C.utf8_stdio()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root")
    ap.add_argument("--out")
    args = ap.parse_args()
    root = Path(args.root).resolve() if args.root else Path(__file__).resolve().parents[2]
    cfg = C.load_toml(root / "tools" / "guard" / "guard.toml")
    out = Path(args.out) if args.out else root / cfg["repo"]["tracker_xlsx"]
    status_path = root / "tracker" / "status.json"
    if not status_path.is_file():
        print("tracker/status.json missing. Run: python tools/tracker/verify.py")
        return 2
    import json
    status = json.loads(status_path.read_text(encoding="utf-8"))
    feats = C.load_yaml(root / "tracker" / "features.yaml")["features"]
    st = status["features"]

    # back up whatever is there (once) if it is not one of ours
    if out.is_file():
        try:
            old = openpyxl.load_workbook(out, read_only=True)
            ours = "_meta" in old.sheetnames
            old.close()
        except Exception:  # noqa: BLE001
            ours = False
        if not ours:
            bak = root / "tracker" / "legacy" / f"replaced-{time.strftime('%Y%m%d-%H%M%S')}-{out.name}"
            bak.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(out, bak)
            print(f"backed up the old workbook to {bak}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feature Tracker"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in feats:
        r = st.get(f["id"])
        if not r:
            continue
        ev = f.get("evidence") or {}
        ws.append([
            f["id"], "; ".join(f.get("spec_refs") or []) or "(not a spec section)", f.get("title"), f.get("category"), f.get("verdict"),
            f.get("source"), f.get("provider"), f.get("layer"), r["label"], num(r["pct"]), num(r["claimed_pct"]), num(r["gap"]),
            r["next"], ", ".join(r["flags"]), ", ".join(str(x) for x in (ev.get("files") or [])), " | ".join(r["reasons"]),
        ])
        row = ws.max_row
        ws.cell(row, 9).fill = PatternFill("solid", fgColor=FILLS.get(r["label"], "FFFFFF"))
        for col in (2, 3, 13, 14, 15, 16):
            ws.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
    last = ws.max_row
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last}"

    # ---------------- Summary: live formulas over the grid, nothing typed by hand
    sm = wb.create_sheet("Summary")
    T = "'Feature Tracker'"
    rng = lambda col: f"{T}!${col}$2:${col}${last}"  # noqa: E731
    rows = [
        ("GENERATED FROM EVIDENCE. Do not edit.", None),
        ("Generated at", status["generated_at"]),
        ("Commit", status["commit"] + (" (dirty)" if status["dirty"] else "")),
        ("Guard verdict", (status.get("guard") or {}).get("verdict", "not run")),
        (None, None),
        ("Features tracked", f"=COUNTA({rng('A')})"),
        ("Verified average %", f"=IFERROR(ROUND(AVERAGE({rng('J')}),1),\"n/a\")"),
        ("OLD claimed average % (unverified, from the previous tracker)", f"=IFERROR(ROUND(AVERAGE({rng('K')}),1),\"n/a\")"),
        ("Features the old tracker claimed >= 25 points above evidence", f"=COUNTIF({rng('L')},\">=25\")"),
        ("Integrated or better (>= 85%)", f"=COUNTIF({rng('J')},\">=85\")"),
        (None, None),
    ]
    for name in ("NO_EVIDENCE", "FILES_PRESENT", "IMPLEMENTED", "UNIT_VERIFIED", "INTEGRATED", "PROVEN", "INHERITED_OK"):
        rows.append((f"Level {name}", f"=COUNTIF({rng('I')},\"{name}\")"))
    sm_spec = status["summary"].get("spec_sections_total")
    rows += [(None, None), ("Spec sections in the spec file", sm_spec)]
    for a, b in rows:
        sm.append([a, b])
    sm["A1"].font = Font(bold=True, color="C00000")
    sm.column_dimensions["A"].width = 70
    sm.column_dimensions["B"].width = 34

    lg = wb.create_sheet("Legend")
    legend = [
        ("Level", "Meaning (a level needs every level below it)", "%"),
        ("NO_EVIDENCE", "Nothing declared, or declared files are missing", 0),
        ("FILES_PRESENT", "Every declared file exists and is non-empty", 15),
        ("IMPLEMENTED", "Declared functions/classes exist and are not stubs", 40),
        ("UNIT_VERIFIED", "Declared unit commands / guard checks pass; tests have real assertions", 65),
        ("INTEGRATED", "Integration test through the Pi path passes", 85),
        ("PROVEN", "Recorded end-to-end proof is current AND a human signed off", 100),
        ("INHERITED_OK", "KEEP-PI feature: file present and packages/ matches the pinned upstream", 100),
        (None, None, None),
        ("Rule", "Nobody types a status. tracker/features.yaml holds evidence; tools/tracker/verify.py computes; this file is regenerated.", None),
    ]
    for row in legend:
        lg.append(list(row))
    lg.column_dimensions["A"].width = 18
    lg.column_dimensions["B"].width = 100

    sc = wb.create_sheet("Spec Coverage")
    sc.append(["Spec line", "Spec section", "Feature ids", "Mapping"])
    spec_path = root / cfg["repo"]["spec_path"]
    if spec_path.is_file():
        heads = C.parse_spec(spec_path)["headings"]
        owners: dict[str, list[str]] = {}
        kind: dict[str, str] = {}
        for f in feats:
            for ref in f.get("spec_refs") or []:
                k = C.norm_key(ref)
                owners.setdefault(k, []).append(f["id"])
                kind[k] = f.get("mapping", "")
        for h in heads:
            ids = owners.get(h["key"], [])
            sc.append([h["line"], C.display_heading(h), ", ".join(ids) or "NOT TRACKED", kind.get(h["key"], "")])
    sc.column_dimensions["B"].width = 70
    sc.column_dimensions["C"].width = 30

    mt = wb.create_sheet("_meta")
    mt.sheet_state = "hidden"
    wb.save(out)

    # the grid hash is taken from the SAVED workbook so it matches what the guard will read back
    reread = openpyxl.load_workbook(out)
    gh = C.grid_hash(reread["Feature Tracker"])
    reread.close()
    wb2 = openpyxl.load_workbook(out)
    mt = wb2["_meta"]
    for k, v in (("grid_sha", gh), ("status_sha", C.sha256_file(status_path)), ("commit", status["commit"]),
                 ("generated_at", status["generated_at"]), ("note", "generated by tools/tracker/export_xlsx.py - hand edits are detected")):
        mt.append([k, v])
    wb2.save(out)
    print(f"wrote {out}  ({last - 1} features)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
