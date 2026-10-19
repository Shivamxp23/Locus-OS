"""Tests for the Locus guard (tools/guard) and the evidence tracker (tools/tracker).

    python -m pip install pytest pyyaml openpyxl
    python -m pytest tests/guard -q

Every test that needs a repository builds a throw-away git repo (a copy of a once-built template), so nothing
here touches your real working tree. Secrets used by the tests are assembled at run time and are fake.
"""
from __future__ import annotations

import importlib.util
import json
import os
import re
import shutil
import stat
import subprocess
import sys
import types
from pathlib import Path

import pytest
import yaml

BUNDLE = Path(__file__).resolve().parents[2]
TOOLS = BUNDLE / "tools"
PY = sys.executable
QPY = f'"{PY}"'


# ================================================================== module loading
def load_tool(name: str, rel: str):
    path = BUNDLE / rel
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="session")
def common():
    sys.path.insert(0, str(TOOLS))
    import locus_common  # noqa: PLC0415
    return locus_common


@pytest.fixture(scope="session")
def check_mod(common):
    return load_tool("locus_check", "tools/guard/check.py")


@pytest.fixture(scope="session")
def verify_mod(common):
    return load_tool("locus_verify", "tools/tracker/verify.py")


# ================================================================== repo helpers
def put(repo: Path, rel: str, text: str = "") -> Path:
    p = repo / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(text)
    return p


def git(repo: Path, *args: str, check: bool = True) -> str:
    p = subprocess.run(["git", *args], cwd=repo, capture_output=True, text=True, encoding="utf-8")
    if check and p.returncode != 0:
        raise AssertionError(f"git {' '.join(args)} failed: {p.stdout}{p.stderr}")
    return p.stdout


def base_env(extra: dict | None = None) -> dict:
    env = {**os.environ, "PYTHONDONTWRITEBYTECODE": "1", "PYTHONIOENCODING": "utf-8"}
    env.pop("LOCUS_ALLOW_GUARD_EDIT", None)
    if extra:
        env.update(extra)
    return env


class Run:
    def __init__(self, rc: int, report: dict, out: str):
        self.rc, self.report, self.out = rc, report, out

    def of(self, check: str, sev: str | None = None) -> list[dict]:
        return [f for f in self.report["findings"] if f["check"] == check and (sev is None or f["severity"] == sev)]

    def fails(self, check: str | None = None) -> list[dict]:
        return [f for f in self.report["findings"] if f["severity"] == "FAIL" and (check is None or f["check"] == check)]

    def failing_checks(self) -> list[str]:
        return sorted({f["check"] for f in self.fails()})

    def where(self, check: str, sev: str = "FAIL") -> list[str]:
        return [f["where"] for f in self.of(check, sev)]


def guard(repo: Path, *args: str, env: dict | None = None) -> Run:
    out = repo / "tracker" / ".cache" / "t_report.json"
    if out.exists():
        out.unlink()
    p = subprocess.run([PY, str(repo / "tools" / "guard" / "check.py"), *args, "--json", str(out)],
                       cwd=repo, capture_output=True, text=True, encoding="utf-8", env=base_env(env))
    assert out.is_file(), f"guard wrote no report. rc={p.returncode}\n{p.stdout}\n{p.stderr}"
    return Run(p.returncode, json.loads(out.read_text(encoding="utf-8")), p.stdout)


def tracker(repo: Path, script: str, *args: str, env: dict | None = None) -> subprocess.CompletedProcess:
    return subprocess.run([PY, str(repo / "tools" / "tracker" / script), *args], cwd=repo, capture_output=True, text=True,
                          encoding="utf-8", env=base_env(env))


def commit_all(repo: Path, msg: str = "wip") -> None:
    git(repo, "add", "-A")
    git(repo, "commit", "-q", "-m", msg)


# ================================================================== the template repo
SPEC = """# Fixture spec

Locus is a meta-harness over other agents. Auditability matters: every action is logged and backed up.
Tasks are prioritised by multi-scale rules. Memory is Neo4j Agent Memory (neo4j-labs/agent-memory).

## Feature §1.1: Alpha tool
**Verdict**: LOCUS-NEW
Lives in `locus-extension/src/tools/alpha.ts`.

## Feature §1.2: Beta engine
**Verdict**: KEEP-PI

## Feature §2.1: Gamma sidecar
**Verdict**: LOCUS-NEW
"""

PI_BRIDGE = '''import shutil
import subprocess


def get_node_executable() -> str:
    exe = shutil.which("node")
    if exe is None:
        raise RuntimeError("node not found")
    return exe


def run_pi_turn(prompt: str, cwd: str) -> str:
    cmd = [get_node_executable(), "pi.js", "--approve", "--no-context-files", "-p", prompt]
    proc = subprocess.run(cmd, cwd=cwd, capture_output=True, text=True)
    return proc.stdout
'''

GITIGNORE = "\n".join([".env", ".venv/", "node_modules/", "dist/", "data/couchdb/*", "!data/couchdb/local.ini",
                       "__pycache__/", "tracker/.cache/", ""])


def build_template(t: Path, common) -> None:
    t.mkdir(parents=True)
    shutil.copytree(TOOLS, t / "tools", ignore=shutil.ignore_patterns("__pycache__", "*.pyc", "upstream.lock.json", "spec_lint.baseline.json"))
    cfg = (TOOLS / "guard" / "guard.toml").read_text(encoding="utf-8")
    cfg = re.sub(r'^spec_path = ".*"$', 'spec_path = "spec/SPEC.md"', cfg, flags=re.M)
    cfg = re.sub(r'^tracker_xlsx = ".*"$', 'tracker_xlsx = "spec/Tracker.xlsx"', cfg, flags=re.M)
    cfg = re.sub(r"allowed_top_level\s*=\s*\[[^\]]*\]", "allowed_top_level = []", cfg, count=1, flags=re.S)
    cfg = re.sub(r"allow_modified\s*=\s*\[[^\]]*\]", "allow_modified = []", cfg, count=1, flags=re.S)
    put(t, "tools/guard/guard.toml", cfg)

    git(t, "init", "-q")
    git(t, "symbolic-ref", "HEAD", "refs/heads/main")
    for k, v in (("user.email", "t@example.com"), ("user.name", "Test"), ("commit.gpgsign", "false"), ("core.autocrlf", "false")):
        git(t, "config", k, v)

    # commit 1: the upstream Pi code
    put(t, "packages/pi-mono/index.ts", "export const upstream = 1;\n")
    put(t, "packages/pi-mono/README.md", "# pi\n")
    git(t, "add", "packages")
    git(t, "commit", "-q", "-m", "pi upstream")

    # commit 2: the Locus overlay
    put(t, ".gitignore", GITIGNORE)
    put(t, "AGENTS.md", "<!-- LOCUS-AGENTS-V1 -->\n# Locus rules\nOne rule source. Keep it short.\n")
    put(t, "CLAUDE.md", "@AGENTS.md\n")
    put(t, "docker-compose.yml", "services:\n  couchdb:\n    image: couchdb:3\n")
    put(t, "sidecar/__init__.py", '__version__ = "0.0.1"\n')
    put(t, "sidecar/pi_bridge.py", PI_BRIDGE)
    put(t, "sidecar/conductor.py", "def tick():\n    return 1\n")
    put(t, "sidecar/life_os.py", "def score():\n    return 1\n")
    put(t, "sidecar/gamma.py", "def gamma(x):\n    y = x * 2\n    return y + 1\n")
    put(t, "locus_cli/__init__.py", "")
    put(t, "locus-extension/src/tools/alpha.ts", "export function alpha(): number {\n  const v = 1;\n  return v;\n}\n")
    put(t, "docs/agent/01-example.md", "# Example segment\nShort and focused.\n")
    for d in ("00-System", "01-Projects", "02-Areas", "03-Resources", "04-Captures", "05-Daily", "06-Content", "07-Archives"):
        put(t, f"vault/{d}/.gitkeep", "")
    put(t, "spec/SPEC.md", SPEC)
    spec = common.parse_spec(t / "spec" / "SPEC.md")
    sha = {common.display_heading(h): h["sha"] for h in spec["headings"]}
    feats = [
        {"id": "F-0001", "title": "Alpha tool", "spec_refs": ["§1.1: Alpha tool"], "mapping": "exact", "verdict": "LOCUS-NEW",
         "layer": "extension", "inherited": False, "manual": False, "claimed": {"pct": 90},
         "spec_sha": {"§1.1: Alpha tool": sha["§1.1: Alpha tool"]},
         "evidence": {"files": ["locus-extension/src/tools/alpha.ts"], "symbols": [], "guard": [], "unit": [], "integration": []}},
        {"id": "F-0002", "title": "Beta engine", "spec_refs": ["§1.2: Beta engine"], "mapping": "exact", "verdict": "KEEP-PI",
         "layer": "pi", "inherited": True, "manual": False, "claimed": {"pct": 100},
         "spec_sha": {"§1.2: Beta engine": sha["§1.2: Beta engine"]},
         "evidence": {"files": ["packages/pi-mono/index.ts"], "symbols": [], "guard": [], "unit": [], "integration": []}},
        {"id": "F-0003", "title": "Gamma sidecar", "spec_refs": ["§2.1: Gamma sidecar"], "mapping": "exact", "verdict": "LOCUS-NEW",
         "layer": "sidecar", "inherited": False, "manual": False, "claimed": {"pct": 50},
         "spec_sha": {"§2.1: Gamma sidecar": sha["§2.1: Gamma sidecar"]},
         "evidence": {"files": ["sidecar/gamma.py"], "symbols": [{"file": "sidecar/gamma.py", "name": "gamma"}],
                      "guard": [], "unit": [], "integration": []}},
    ]
    put(t, "tracker/features.yaml", yaml.safe_dump({"schema": 1, "spec_file": "spec/SPEC.md", "features": feats}, sort_keys=False, allow_unicode=True))
    put(t, "tracker/spec_exclusions.yaml", "exclusions: []\n")
    put(t, "tracker/overrides.yaml", "{}\n")
    put(t, "tracker/signoff.yaml", "{}\n")
    commit_all(t, "locus overlay")

    env = base_env()
    for cmd in ([PY, "tools/guard/pin_upstream.py"], [PY, "tools/guard/check.py", "--adopt-top-level"],
                [PY, "tools/guard/check.py", "--update-spec-baseline"]):
        p = subprocess.run(cmd, cwd=t, capture_output=True, text=True, encoding="utf-8", env=env)
        assert p.returncode == 0, f"{cmd}\n{p.stdout}\n{p.stderr}"
    commit_all(t, "pin upstream, adopt layout, baseline spec")


@pytest.fixture(scope="session")
def template(tmp_path_factory, common) -> Path:
    t = tmp_path_factory.mktemp("tpl") / "repo"
    build_template(t, common)
    return t


@pytest.fixture
def repo(template, tmp_path) -> Path:
    dst = tmp_path / "repo"
    shutil.copytree(template, dst, symlinks=True)
    return dst


# ================================================================== 1. the baseline must be clean
def test_template_passes_full_guard(repo):
    r = guard(repo, "--full")
    assert r.rc == 0, r.out
    assert r.fails() == []


def test_fast_mode_also_passes_and_skips_full_checks(repo):
    r = guard(repo, "--fast")
    assert r.rc == 0, r.out
    assert r.of("spec-lint") == [] and r.of("upstream-clean") == []


def test_default_report_location_and_json_flag(repo):
    p = subprocess.run([PY, str(repo / "tools" / "guard" / "check.py"), "--fast"], cwd=repo, capture_output=True, text=True,
                       encoding="utf-8", env=base_env())
    assert p.returncode == 0, p.stdout
    rep = json.loads((repo / "tracker" / ".cache" / "guard_report.json").read_text(encoding="utf-8"))
    assert rep["verdict"] == "PASS" and "findings" in rep


def test_strict_mode_turns_warnings_into_failures(repo):
    r = guard(repo, "--full", "--strict")  # no git remote in the fixture -> WARN git
    assert r.rc == 1 and r.of("git", "WARN")


def test_only_filter_runs_just_the_named_check(repo):
    os_dir = repo / "core"
    os_dir.mkdir()
    (os_dir / "x.py").write_text("x = 1\n", encoding="utf-8")
    r = guard(repo, "--only", "layout")
    assert r.of("forbidden-path") == []  # not selected
    r2 = guard(repo, "--only", "forbidden-paths")
    assert r2.of("forbidden-path", "FAIL")


def test_missing_config_exits_2(repo):
    p = subprocess.run([PY, str(repo / "tools" / "guard" / "check.py"), "--config", str(repo / "nope.toml")], cwd=repo,
                       capture_output=True, text=True, encoding="utf-8", env=base_env())
    assert p.returncode == 2


def test_a_crashing_check_is_reported_as_fail_not_ignored(repo):
    cfg = repo / "tools" / "guard" / "guard.toml"
    cfg.write_text(cfg.read_text(encoding="utf-8") + '\n[[forbidden_text]]\nid = "bad"\npattern = "("\nglobs = ["*.py"]\nmessage = "x"\n', encoding="utf-8")
    r = guard(repo, "--fast")
    assert r.rc == 1
    assert any("crashed" in f["message"] for f in r.fails())


# ================================================================== 2. the exact mix-ups that already happened
def test_core_folder_is_forbidden(repo):
    put(repo, "core/pi_bridge.py", "x = 1\n")
    r = guard(repo, "--fast")
    assert r.rc == 1
    assert "core" in r.where("forbidden-path")


def test_nested_sidecar_core_is_forbidden(repo):
    put(repo, "sidecar/core/x.py", "x = 1\n")
    assert "sidecar/core" in guard(repo, "--fast").where("forbidden-path")


@pytest.mark.parametrize("line", ["from core.pi_bridge import run_pi_turn", "import core", "  from core import conductor", "import core.api as api"])
def test_import_of_core_fails(repo, line):
    put(repo, "sidecar/uses_core.py", line + "\n")
    r = guard(repo, "--fast")
    assert r.of("py-import-core", "FAIL"), r.out


@pytest.mark.parametrize("line", ["from sidecar.core_utils import x", "import corelib", "from encore import x", "# from core import x",
                                  "from sidecar import core_thing"])
def test_lookalike_imports_are_not_flagged(repo, line):
    put(repo, "sidecar/ok_imports.py", line + "\n")
    assert guard(repo, "--only", "forbidden-text").of("py-import-core") == []


@pytest.mark.parametrize("line", ['sys.modules["core"] = sidecar', "sys.modules['core']=sidecar", "sys.modules.setdefault('core', sidecar)"])
def test_core_alias_shim_fails(repo, line):
    put(repo, "sidecar/shim.py", "import sys\nimport sidecar\n" + line + "\n")
    assert guard(repo, "--fast").of("core-alias-shim", "FAIL")


def test_other_sys_modules_use_is_fine(repo):
    put(repo, "sidecar/ok.py", 'import sys\nm = sys.modules["sidecar"]\n')
    assert guard(repo, "--only", "forbidden-text").of("core-alias-shim") == []


@pytest.mark.parametrize("name,text", [
    ("run.ps1", "python -m core.conductor\n"),
    ("start.sh", "uvicorn core.api:app --port 1\n"),
    ("docker-compose.override.yml", "services:\n  s:\n    command: gunicorn core\n"),
])
def test_launching_core_fails(repo, name, text):
    put(repo, name, text)
    assert guard(repo, "--fast").of("core-module-run", "FAIL")


@pytest.mark.parametrize("name,text", [
    ("go.sh", "cd core/\n"),
    ("go.ps1", 'python "core/main.py"\n'),
    ("go.bat", ".\\core\\run.bat\n"),
    ("svc.service", "ExecStart=/usr/bin/python =core/x.py\n"),
])
def test_core_paths_in_scripts_fail(repo, name, text):
    put(repo, name, text)
    assert guard(repo, "--fast").of("core-path-in-scripts", "FAIL"), name


@pytest.mark.parametrize("text", ["cd sidecar/core_utils/\n", "cd hardcore/\n", "python sidecar/core.py\n", "echo encore/x\n"])
def test_lookalike_script_paths_are_not_flagged(repo, text):
    put(repo, "ok.sh", text)
    assert guard(repo, "--only", "forbidden-text").of("core-path-in-scripts") == []


def test_symlinked_directory_is_a_reparse_point(repo):
    try:
        os.symlink(repo / "sidecar", repo / "alias_dir", target_is_directory=True)
    except (OSError, NotImplementedError):
        pytest.skip("cannot create symlinks here")
    r = guard(repo, "--fast")
    assert r.rc == 1 and "alias_dir" in r.where("reparse-point")


def test_symlinked_file_is_flagged(repo):
    try:
        os.symlink(repo / "AGENTS.md", repo / "sidecar" / "link.md")
    except (OSError, NotImplementedError):
        pytest.skip("cannot create symlinks here")
    assert "sidecar/link.md" in guard(repo, "--fast").where("reparse-point")


def test_reparse_allowlist_is_honoured(repo):
    try:
        os.symlink(repo / "sidecar", repo / "alias_dir", target_is_directory=True)
    except (OSError, NotImplementedError):
        pytest.skip("cannot create symlinks here")
    cfg = repo / "tools" / "guard" / "guard.toml"
    cfg.write_text(cfg.read_text(encoding="utf-8").replace("allow = []", 'allow = ["alias_dir"]', 1), encoding="utf-8")
    assert guard(repo, "--only", "reparse").of("reparse-point") == []


def test_windows_junction_and_symlink_tags_are_detected_but_onedrive_is_not(check_mod, monkeypatch, tmp_path):
    def fake(tag):
        return types.SimpleNamespace(st_mode=stat.S_IFDIR | 0o755, st_reparse_tag=tag)
    p = tmp_path / "x"
    p.mkdir()
    for tag, expect in ((0xA0000003, "junction"), (0xA000000C, "symlink"), (0x9000701A, None), (0, None)):
        monkeypatch.setattr(os, "lstat", lambda _p, t=tag: fake(t))
        assert check_mod.reparse_kind(p) == expect, hex(tag)


# ================================================================== 3. layout and top-level discipline
def test_missing_required_folder_and_file_fail(repo):
    shutil.rmtree(repo / "locus_cli")
    (repo / "AGENTS.md").unlink()
    r = guard(repo, "--fast")
    assert set(r.where("layout-required")) == {"locus_cli", "AGENTS.md"}


def test_new_top_level_entry_fails_after_adoption(repo):
    put(repo, "scratch/notes.txt", "x\n")
    r = guard(repo, "--fast")
    assert "scratch" in r.where("layout-top-level")


def test_top_level_file_also_fails(repo):
    put(repo, "random_script.py", "x = 1\n")
    assert "random_script.py" in guard(repo, "--fast").where("layout-top-level")


def test_adoption_flow(repo):
    cfg = repo / "tools" / "guard" / "guard.toml"
    cfg.write_text(re.sub(r"allowed_top_level\s*=\s*\[[^\]]*\]", "allowed_top_level = []", cfg.read_text(encoding="utf-8")), encoding="utf-8")
    put(repo, "scratch/notes.txt", "x\n")
    r = guard(repo, "--fast")
    assert r.of("layout-baseline", "WARN") and "scratch" in r.of("layout-baseline")[0]["message"]
    p = subprocess.run([PY, str(repo / "tools" / "guard" / "check.py"), "--adopt-top-level"], cwd=repo, capture_output=True, text=True,
                       encoding="utf-8", env=base_env())
    assert p.returncode == 0
    text = cfg.read_text(encoding="utf-8")
    assert '"scratch"' in text and '"spec"' in text and '"vault"' in text
    assert guard(repo, "--fast").of("layout-top-level") == []
    put(repo, "another/x.txt", "x\n")
    assert "another" in guard(repo, "--fast").where("layout-top-level")


def test_adoption_refuses_while_a_forbidden_folder_exists(repo):
    put(repo, "core/x.py", "x = 1\n")
    p = subprocess.run([PY, str(repo / "tools" / "guard" / "check.py"), "--adopt-top-level"], cwd=repo, capture_output=True, text=True,
                       encoding="utf-8", env=base_env())
    assert p.returncode == 2 and "Refusing" in p.stdout


# ================================================================== 4. vault naming
def test_noncanonical_vault_folder_fails_with_a_hint(repo):
    put(repo, "vault/06-Resources/note.md", "x\n")
    r = guard(repo, "--fast")
    f = r.of("vault-layout", "FAIL")
    assert f and "03-Resources" in f[0]["message"]


def test_hidden_vault_folders_are_ignored(repo):
    put(repo, "vault/.obsidian/app.json", "{}\n")
    assert guard(repo, "--fast").of("vault-layout") == []


@pytest.mark.parametrize("name,hint", [("05-Logs", "05-Daily/Logs"), ("07-Diagrams", "03-Resources/Diagrams"), ("00-System-Data", "00-System"),
                                       ("01-Daily", "05-Daily"), ("02-Reviews", "05-Daily/Weekly")])
def test_code_using_old_vault_names_fails(repo, name, hint):
    put(repo, "sidecar/paths.py", f'VAULT_DIR = "vault/{name}/x"\n')
    f = guard(repo, "--fast").of("vault-name", "FAIL")
    assert f and hint in f[0]["message"]


def test_canonical_vault_names_and_dates_are_fine(repo):
    put(repo, "sidecar/paths.py", 'A = "vault/04-Captures/x"\nB = "2026-10-04T05:06"\nC = "2024-01-15"\nD = "v01-alpha"\n')
    assert guard(repo, "--fast").of("vault-name") == []


# ================================================================== 5. upstream Pi stays untouched
def test_modifying_a_pi_file_fails(repo):
    put(repo, "packages/pi-mono/index.ts", "export const upstream = 2;\n")
    r = guard(repo, "--full")
    assert "packages/pi-mono/index.ts" in r.where("upstream-clean")


def test_adding_a_file_inside_packages_fails(repo):
    put(repo, "packages/pi-mono/locus-hack.ts", "export {};\n")
    assert "packages/pi-mono/locus-hack.ts" in guard(repo, "--full").where("upstream-clean")


def test_committed_edits_to_pi_are_still_caught(repo):
    put(repo, "packages/pi-mono/index.ts", "export const upstream = 3;\n")
    commit_all(repo, "oops")
    assert "packages/pi-mono/index.ts" in guard(repo, "--full").where("upstream-clean")


def test_allow_modified_is_a_human_decision_in_guard_toml(repo):
    put(repo, "packages/pi-mono/index.ts", "export const upstream = 2;\n")
    cfg = repo / "tools" / "guard" / "guard.toml"
    cfg.write_text(cfg.read_text(encoding="utf-8").replace("allow_modified = []", 'allow_modified = ["packages/pi-mono/index.ts"]'), encoding="utf-8")
    assert guard(repo, "--full").of("upstream-clean", "FAIL") == []


def test_missing_pin_fails(repo):
    (repo / "tools" / "guard" / "upstream.lock.json").unlink()
    assert guard(repo, "--full").of("upstream-clean", "FAIL")


def test_pin_pointing_nowhere_fails(repo):
    lock = repo / "tools" / "guard" / "upstream.lock.json"
    d = json.loads(lock.read_text(encoding="utf-8"))
    d["baseline_commit"] = "0" * 40
    lock.write_text(json.dumps(d), encoding="utf-8")
    assert "not in this repo" in guard(repo, "--full").of("upstream-clean", "FAIL")[0]["message"]


def test_pin_upstream_picks_the_commit_before_locus_started(repo):
    first = git(repo, "rev-list", "--max-parents=0", "HEAD").strip()
    lock = json.loads((repo / "tools" / "guard" / "upstream.lock.json").read_text(encoding="utf-8"))
    assert lock["baseline_commit"] == first


def test_pin_upstream_compare_dir(repo, tmp_path):
    ref = tmp_path / "fresh"
    shutil.copytree(repo / "packages", ref / "packages")
    run = lambda: subprocess.run([PY, str(repo / "tools" / "guard" / "pin_upstream.py"), "--compare-dir", str(ref)], cwd=repo,  # noqa: E731
                                 capture_output=True, text=True, encoding="utf-8", env=base_env())
    assert run().returncode == 0
    put(repo, "packages/pi-mono/index.ts", "changed\n")
    p = run()
    assert p.returncode == 1 and "MODIFIED" in p.stdout


# ================================================================== 6. a re-export must not replace the real code
def test_reexport_replacing_the_implementation_fails(repo):
    put(repo, "sidecar/pi_bridge.py", "from sidecar.other import run_pi_turn, get_node_executable\n# --approve --no-context-files\n")
    r = guard(repo, "--fast")
    msgs = [f["message"] for f in r.of("must-define", "FAIL")]
    assert len(msgs) == 2 and all("only imported" in m for m in msgs)


def test_missing_function_and_missing_file_fail(repo):
    put(repo, "sidecar/pi_bridge.py", 'def get_node_executable():\n    a = 1\n    b = 2\n    return "--approve --no-context-files"\n')
    assert any("run_pi_turn" in f["message"] for f in guard(repo, "--fast").of("must-define", "FAIL"))
    (repo / "sidecar" / "pi_bridge.py").unlink()
    assert any("missing" in f["message"] for f in guard(repo, "--fast").of("must-define", "FAIL"))


def test_stub_sized_body_warns(repo):
    put(repo, "sidecar/pi_bridge.py", 'def get_node_executable():\n    return "node --approve --no-context-files"\n\n\ndef run_pi_turn(p, c):\n    return p\n')
    assert guard(repo, "--fast").of("must-define", "WARN")


def test_bridge_must_pass_the_flags_to_pi(repo):
    put(repo, "sidecar/pi_bridge.py", PI_BRIDGE.replace('"--approve", ', ""))
    assert any("--approve" in f["message"] for f in guard(repo, "--fast").of("pi-bridge-flags", "FAIL"))
    put(repo, "sidecar/pi_bridge.py", PI_BRIDGE.replace('"--no-context-files", ', ""))
    assert guard(repo, "--fast").of("pi-bridge-flags", "FAIL")
    put(repo, "sidecar/pi_bridge.py", PI_BRIDGE.replace('"--no-context-files"', '"-nc"'))  # the short flag is accepted
    assert guard(repo, "--fast").of("pi-bridge-flags") == []


# ================================================================== 7. python syntax, secrets, agent files
def test_syntax_error_in_a_package_fails(repo):
    put(repo, "sidecar/bad.py", "def (:\n")
    assert any(w.startswith("sidecar/bad.py") for w in guard(repo, "--fast").where("python-syntax"))


def fake_secrets() -> dict[str, str]:
    a = "A" * 40
    return {
        "groq": "gsk_" + a[:24],
        "openrouter": "sk-or-v1-" + a[:24],
        "anthropic": "sk-ant-" + a,
        "google": "AIza" + "B" * 35,
        "telegram": "123456789:AA" + "C" * 33,
        "github": "ghp_" + "D" * 36,
        "slack": "xoxb-" + "1" * 12,
        "private-key": "-----BEGIN " + "PRIVATE KEY-----",
        "couchdb-default": "locus_admin" + "_secret",
    }


@pytest.mark.parametrize("kind", sorted(fake_secrets()))
def test_secrets_are_detected(repo, kind):
    put(repo, "sidecar/config_oops.py", f'KEY = "{fake_secrets()[kind]}"\n')
    r = guard(repo, "--fast")
    assert r.rc == 1 and any("config_oops.py" in w for w in r.where("secrets")), r.out


def test_secret_in_a_staged_file_is_caught_by_the_hook_mode(repo):
    put(repo, "sidecar/config_oops.py", f'KEY = "{fake_secrets()["groq"]}"\n')
    git(repo, "add", "sidecar/config_oops.py")
    assert guard(repo, "--fast", "--hook").of("secrets", "FAIL")


def test_secret_allow_marker_and_allowed_files(repo):
    put(repo, "sidecar/doc_example.py", f'# guard:allow-secret (documentation example)\nKEY = "{fake_secrets()["groq"]}"\n')
    put(repo, ".env.example", f"GROQ_API_KEY={fake_secrets()['groq']}\n")
    assert guard(repo, "--fast").of("secrets") == []


def test_an_env_file_that_git_could_commit_fails(repo):
    put(repo, ".gitignore", "node_modules/\n")  # .env is no longer ignored
    put(repo, ".env", "GROQ_API_KEY=abc\n")
    assert ".env" in guard(repo, "--fast").where("secrets")
    (repo / ".env").unlink()
    put(repo, "sidecar/.env.local", "A=1\n")
    assert guard(repo, "--fast").of("secrets", "FAIL")


def test_an_ignored_env_file_is_not_a_commit_risk(repo):
    put(repo, ".env", "GROQ_API_KEY=abc\n")  # the fixture .gitignore lists .env
    assert guard(repo, "--fast").of("secrets") == []


def test_short_lookalikes_are_not_secrets(repo):
    put(repo, "sidecar/ok.py", 'A = "gsk_short"\nB = "task-ant-xyz"\nC = "AIza123"\n')
    assert guard(repo, "--fast").of("secrets") == []


def test_agents_md_without_marker_fails(repo):
    put(repo, "AGENTS.md", "# Development Rules (Pi upstream)\nLong upstream text\n")
    f = guard(repo, "--fast").of("agents-md", "FAIL")
    assert f and "LOCUS-AGENTS-V1" in f[0]["message"]


def test_agents_md_size_limit(repo):
    put(repo, "AGENTS.md", "<!-- LOCUS-AGENTS-V1 -->\n" + "rule rule rule rule\n" * 700)
    assert any("bytes" in f["message"] for f in guard(repo, "--fast").of("agents-md", "FAIL"))


def test_claude_md_must_import_agents_md(repo):
    put(repo, "CLAUDE.md", "# separate rules\n")
    assert guard(repo, "--fast").of("claude-md", "FAIL")
    (repo / "CLAUDE.md").unlink()
    assert guard(repo, "--fast").of("claude-md", "WARN")


def test_root_gemini_md_is_forbidden(repo):
    put(repo, "GEMINI.md", "other rules\n")
    assert "GEMINI.md" in guard(repo, "--fast").where("gemini-md")


def test_oversized_segment_warns_but_errata_and_prompts_are_exempt(repo):
    big = "# big\n" + "line of text that keeps going\n" * 400
    put(repo, "docs/agent/05-big.md", big)
    put(repo, "docs/agent/09-spec-errata.md", big)
    put(repo, "docs/agent/prompts/P1.md", big)
    where = guard(repo, "--fast").where("segment-size", "WARN")
    assert where == ["docs/agent/05-big.md"]


def test_gitignore_warns_when_a_required_line_is_missing(repo):
    put(repo, ".gitignore", ".env\n")
    f = guard(repo, "--full").of("gitignore", "WARN")
    assert f and "node_modules/" in f[0]["message"]


# ================================================================== 8. human review gate
def test_changes_to_rule_files_need_a_human_in_hook_mode(repo):
    put(repo, "AGENTS.md", "<!-- LOCUS-AGENTS-V1 -->\n# edited by an agent\n")
    git(repo, "add", "AGENTS.md")
    r = guard(repo, "--fast", "--hook")
    assert r.rc == 1 and r.of("human-review", "FAIL")
    r2 = guard(repo, "--fast", "--hook", env={"LOCUS_ALLOW_GUARD_EDIT": "1"})
    assert r2.of("human-review", "FAIL") == [] and r2.of("human-review", "WARN")


@pytest.mark.parametrize("rel", ["tools/guard/guard.toml", "tools/tracker/verify.py", "tools/tracker/backlog.py", "tracker/signoff.yaml", "tracker/backlog.yaml", "tracker/sprints.jsonl", "tracker/retro/S1.md", "docs/agent/01-example.md", ".githooks/pre-commit"])
def test_every_protected_path_triggers_the_gate(repo, rel):
    put(repo, rel, (repo / rel).read_text(encoding="utf-8") + "\n# touched\n" if (repo / rel).exists() else "# new\n")
    git(repo, "add", "-f", rel)
    assert guard(repo, "--fast", "--hook").of("human-review", "FAIL"), rel


def test_deleting_a_protected_file_is_also_gated(repo):
    git(repo, "rm", "-q", "tools/tracker/show.py")
    assert guard(repo, "--fast", "--hook").of("human-review", "FAIL")


def test_ordinary_code_changes_do_not_trigger_the_gate(repo):
    put(repo, "sidecar/gamma.py", "def gamma(x):\n    y = x * 3\n    return y + 1\n")
    git(repo, "add", "sidecar/gamma.py")
    assert guard(repo, "--fast", "--hook").of("human-review") == []


def test_outside_hook_mode_the_gate_only_warns(repo):
    put(repo, "AGENTS.md", "<!-- LOCUS-AGENTS-V1 -->\n# changed\n")
    r = guard(repo, "--full")
    assert r.of("human-review", "WARN") and r.of("human-review", "FAIL") == []


# ================================================================== 9. spec lint
def spec_issues(check_mod, common, repo: Path) -> dict[str, str]:
    cfg = common.load_toml(repo / "tools" / "guard" / "guard.toml")
    return check_mod.spec_issues(check_mod.Ctx(repo, cfg, False, False, False))


def test_clean_fixture_spec_has_no_defects(check_mod, common, repo):
    assert spec_issues(check_mod, common, repo) == {}


def test_each_kind_of_spec_defect_is_detected(check_mod, common, repo):
    spec = repo / "spec" / "SPEC.md"
    original = spec.read_text(encoding="utf-8")
    cases = {
        "dup-id:1.1": original + "\n## Feature §1.1: Another alpha\ntext\n",
        "tws-formula-conflict": original + "\nTWS = (a + b)\nTWS = (c * d)\n",
        "vault-name:06-Resources": original + "\nSaved to `06-Resources/`\n",
        "stale-path:stale-locus-core": original + "\nSee `locus/core/x.ts`\n",
        "principle-missing:meta-harness": original.replace("meta-harness", "harness"),
        "feature-count-claim:85": original + "\n85 features total\n",
    }
    for key, text in cases.items():
        spec.write_text(text, encoding="utf-8")
        assert key in spec_issues(check_mod, common, repo), key
    spec.write_text(original, encoding="utf-8")


def test_index_table_conflicts_need_genuinely_different_names(check_mod, common, repo):
    spec = repo / "spec" / "SPEC.md"
    head = "| 4.1 | Foo Bar |\n| 4.1 | Totally Different |\n| 4.2 | Foo Bar |\n| 4.2 | Foo Bar Baz |\n\n"
    spec.write_text(head + spec.read_text(encoding="utf-8"), encoding="utf-8")
    issues = spec_issues(check_mod, common, repo)
    assert "table-conflict:4.1" in issues and "table-conflict:4.2" not in issues


def test_new_spec_defect_fails_known_one_only_warns_and_fixed_one_is_info(repo):
    spec = repo / "spec" / "SPEC.md"
    original = spec.read_text(encoding="utf-8")
    spec.write_text(original + "\nTWS = (a + b)\nTWS = (c * d)\n", encoding="utf-8")
    r = guard(repo, "--full")
    assert any("tws-formula-conflict" in f["where"] for f in r.of("spec-lint", "FAIL"))
    p = subprocess.run([PY, str(repo / "tools" / "guard" / "check.py"), "--update-spec-baseline"], cwd=repo, capture_output=True,
                       text=True, encoding="utf-8", env=base_env())
    assert p.returncode == 0
    r2 = guard(repo, "--full")
    assert r2.of("spec-lint", "FAIL") == [] and r2.of("spec-lint", "WARN")
    spec.write_text(original, encoding="utf-8")
    r3 = guard(repo, "--full")
    assert any("resolved" in f["message"] for f in r3.of("spec-lint", "INFO"))


# ================================================================== 10. tracker integrity (guard side)
def edit_features(repo: Path, fn) -> None:
    p = repo / "tracker" / "features.yaml"
    data = yaml.safe_load(p.read_text(encoding="utf-8"))
    fn(data["features"])
    p.write_text(yaml.safe_dump(data, sort_keys=False, allow_unicode=True), encoding="utf-8")


def test_status_keys_in_features_yaml_are_rejected(repo):
    edit_features(repo, lambda fs: fs[0].update({"status": "complete", "pct": 100}))
    r = guard(repo, "--full")
    assert r.rc == 1 and any("computed from evidence" in f["message"] for f in r.of("tracker", "FAIL"))


def test_duplicate_ids_and_escaping_evidence_paths_fail(repo):
    def dup(fs):
        fs.append(dict(fs[0]))
        fs[1]["evidence"] = {"files": ["../outside.py"]}
    edit_features(repo, dup)
    msgs = " | ".join(f["message"] for f in guard(repo, "--full").of("tracker", "FAIL"))
    assert "Duplicate feature id" in msgs and "escapes the repo" in msgs


def test_spec_reference_that_does_not_exist_fails(repo):
    edit_features(repo, lambda fs: fs[0].update({"spec_refs": ["§9.9: Nothing"]}))
    assert any("spec ref not found" in f["message"] for f in guard(repo, "--full").of("tracker", "FAIL"))


def test_spec_section_without_a_feature_fails_until_excluded(repo):
    spec = repo / "spec" / "SPEC.md"
    spec.write_text(spec.read_text(encoding="utf-8") + "\n## Feature §3.1: Delta\n**Verdict**: LOCUS-NEW\n", encoding="utf-8")
    subprocess.run([PY, str(repo / "tools" / "guard" / "check.py"), "--update-spec-baseline"], cwd=repo, env=base_env(), capture_output=True)
    assert any("Delta" in f["message"] for f in guard(repo, "--full").of("tracker", "FAIL"))
    put(repo, "tracker/spec_exclusions.yaml", 'exclusions:\n  - ref: "§3.1: Delta"\n    reason: "deferred"\n')
    assert guard(repo, "--full").of("tracker", "FAIL") == []


# ================================================================== 11. verify.py: the evidence ladder
def feat(fid="F-9001", files=None, symbols=None, unit=None, integration=None, guard_ids=None, **kw) -> dict:
    f = {"id": fid, "title": "T", "evidence": {"files": files or [], "symbols": symbols or [], "unit": unit or [],
                                                 "integration": integration or [], "guard": guard_ids or []}}
    f.update(kw)
    return f


def ev(verify_mod, repo, f, *, run=True, guard_report=None, upstream_ok=True, signoff=None, overrides=None, spec_sha=None, ids=None):
    return verify_mod.evaluate(repo, f, run, guard_report, upstream_ok, repo / "tracker" / "proof", signoff or {}, overrides or {},
                               spec_sha or {}, ids)


ALPHA_PY = "def compute(a, b):\n    total = a + b\n    return total\n"
GOOD_TEST = "from sidecar.alpha import compute\n\n\ndef test_compute_adds():\n    assert compute(2, 3) == 5\n"


def pytest_cmd(path: str) -> str:
    return f"{QPY} -m pytest {path} -q -p no:cacheprovider"


@pytest.fixture
def alpha(repo):
    put(repo, "sidecar/alpha.py", ALPHA_PY)
    put(repo, "tests/unit/test_alpha.py", GOOD_TEST)
    put(repo, "tests/integration/test_alpha_flow.py",
        "from sidecar.alpha import compute\n\n\ndef test_flow():\n    assert compute(10, 5) == 15\n    assert compute(0, 0) == 0\n")
    return repo


def alpha_feature(**kw) -> dict:
    """A feature whose evidence is real; keyword args named like evidence fields replace them, the rest are feature fields."""
    over = {k: kw.pop(k) for k in ("files", "symbols", "unit", "integration", "guard") if k in kw}
    f = feat(files=["sidecar/alpha.py"], symbols=[{"file": "sidecar/alpha.py", "name": "compute"}],
             unit=[pytest_cmd("tests/unit/test_alpha.py")], integration=[pytest_cmd("tests/integration/test_alpha_flow.py")])
    f["evidence"].update(over)
    f.update(kw)
    return f


def test_no_evidence_is_level_0(verify_mod, repo):
    r = ev(verify_mod, repo, feat(claimed={"pct": 80}))
    assert (r["label"], r["pct"]) == ("NO_EVIDENCE", 0) and "NO_FILES_NAMED" in r["flags"]


def test_missing_files_are_level_0_and_flagged(verify_mod, repo):
    r = ev(verify_mod, repo, feat(files=["sidecar/nope.py"]))
    assert r["label"] == "NO_EVIDENCE" and "CLAIMED_FILES_MISSING" in r["flags"] and "nope.py" in r["next"]


def test_empty_file_does_not_count_as_present(verify_mod, repo):
    put(repo, "sidecar/empty.py", "")
    assert ev(verify_mod, repo, feat(files=["sidecar/empty.py"]))["level"] == 0


def test_files_only_is_level_1(verify_mod, repo):
    r = ev(verify_mod, repo, feat(files=["sidecar/gamma.py"]))
    assert (r["label"], r["pct"]) == ("FILES_PRESENT", 15) and any("no symbols" in x for x in r["reasons"])


def test_directory_and_glob_evidence(verify_mod, repo):
    assert ev(verify_mod, repo, feat(files=["sidecar/"]))["level"] == 1
    assert ev(verify_mod, repo, feat(files=["sidecar/*.py"]))["level"] == 1
    assert ev(verify_mod, repo, feat(files=["sidecar/*.rs"]))["level"] == 0


def test_symbols_make_level_2(verify_mod, repo):
    r = ev(verify_mod, repo, feat(files=["sidecar/gamma.py"], symbols=[{"file": "sidecar/gamma.py", "name": "gamma"}]))
    assert (r["label"], r["pct"]) == ("IMPLEMENTED", 40)
    assert "no unit tests" in r["next"] or "evidence.unit" in r["next"]


@pytest.mark.parametrize("body", ["def f():\n    pass\n", "def f():\n    ...\n", 'def f():\n    """doc"""\n', "def f():\n    raise NotImplementedError\n",
                                  "def f():\n    raise NotImplementedError('later')\n"])
def test_stub_functions_do_not_count(verify_mod, repo, body):
    put(repo, "sidecar/stubby.py", body)
    r = ev(verify_mod, repo, feat(files=["sidecar/stubby.py"], symbols=[{"file": "sidecar/stubby.py", "name": "f"}]))
    assert r["level"] == 1 and any("stub" in x for x in r["reasons"])


def test_class_with_only_stub_methods_is_a_stub(verify_mod, repo):
    put(repo, "sidecar/stubby.py", "class C:\n    def a(self):\n        pass\n    def b(self):\n        ...\n")
    r = ev(verify_mod, repo, feat(files=["sidecar/stubby.py"], symbols=[{"file": "sidecar/stubby.py", "name": "C"}]))
    assert r["level"] == 1


def test_undefined_symbol_and_missing_substring_and_bad_pattern(verify_mod, repo):
    f = feat(files=["sidecar/gamma.py"], symbols=[{"file": "sidecar/gamma.py", "name": "delta", "contains": ["nope"], "pattern": "zzz+"}])
    r = ev(verify_mod, repo, f)
    text = " | ".join(r["reasons"])
    assert r["level"] == 1 and "`delta` is not defined" in text and "`nope`" in text and "zzz" in text


def test_stub_marker_blocks_level_2(verify_mod, repo):
    put(repo, "sidecar/marked.py", "def real(x):\n    # STUB(F-9001) replace with the real thing\n    y = x\n    return y\n")
    r = ev(verify_mod, repo, feat(files=["sidecar/marked.py"], symbols=[{"file": "sidecar/marked.py", "name": "real"}]))
    assert r["level"] == 1 and any("STUB(F-9001)" in x for x in r["reasons"])


def test_a_comment_cannot_satisfy_a_contains_check_python(verify_mod, repo):
    put(repo, "sidecar/c.py", "# locus_checkin is registered below\nREGISTRY = []\n")
    r = ev(verify_mod, repo, feat(files=["sidecar/c.py"], symbols=[{"file": "sidecar/c.py", "contains": ["locus_checkin"]}]))
    assert r["level"] == 1
    put(repo, "sidecar/c.py", 'REGISTRY = ["locus_checkin"]  # fine\n')
    assert ev(verify_mod, repo, feat(files=["sidecar/c.py"], symbols=[{"file": "sidecar/c.py", "contains": ["locus_checkin"]}]))["level"] == 2


def test_a_comment_cannot_satisfy_a_contains_check_typescript(verify_mod, repo):
    put(repo, "locus-extension/src/tools/index.ts",
        "// registerTool('locus_checkin')\n/* locus_life */\nexport const tools = [\"locus_goals\"]; // trailing locus_memory\n")
    f = feat(files=["locus-extension/src/tools/index.ts"],
             symbols=[{"file": "locus-extension/src/tools/index.ts", "contains": ["locus_goals"]}])
    assert ev(verify_mod, repo, f)["level"] == 2
    for hidden in ("locus_checkin", "locus_life", "locus_memory"):
        f["evidence"]["symbols"][0]["contains"] = [hidden]
        assert ev(verify_mod, repo, f)["level"] == 1, hidden


def test_typescript_symbols_and_stubs(verify_mod, repo):
    put(repo, "locus-extension/src/real.ts", "export function registerReal(pi: any) {\n  pi.registerTool({ name: 'x' });\n  return 1;\n}\n")
    put(repo, "locus-extension/src/fake.ts", "export function registerFake(pi: any) {\n  throw new Error('not implemented');\n}\n")
    put(repo, "locus-extension/src/arrow.ts", "export const arrowTool = async (pi: any) => {\n  return 1;\n};\n")
    ok = feat(files=["locus-extension/src/real.ts"], symbols=[{"file": "locus-extension/src/real.ts", "name": "registerReal"}])
    bad = feat(files=["locus-extension/src/fake.ts"], symbols=[{"file": "locus-extension/src/fake.ts", "name": "registerFake"}])
    arrow = feat(files=["locus-extension/src/arrow.ts"], symbols=[{"file": "locus-extension/src/arrow.ts", "name": "arrowTool"}])
    assert ev(verify_mod, repo, ok)["level"] == 2
    assert ev(verify_mod, repo, bad)["level"] == 1
    assert ev(verify_mod, repo, arrow)["level"] == 2


def test_passing_real_unit_test_reaches_level_3(verify_mod, alpha):
    f = alpha_feature(integration=[])
    r = ev(verify_mod, alpha, f)
    assert (r["label"], r["pct"]) == ("UNIT_VERIFIED", 65), r["reasons"]


def test_static_mode_never_credits_a_test_it_did_not_run(verify_mod, alpha):
    r = ev(verify_mod, alpha, alpha_feature(integration=[]), run=False)
    assert r["level"] == 2 and any("not run" in x for x in r["reasons"])


def test_failing_unit_test_stays_level_2(verify_mod, alpha):
    put(alpha, "tests/unit/test_alpha.py", "from sidecar.alpha import compute\n\n\ndef test_compute_adds():\n    assert compute(2, 3) == 6\n")
    r = ev(verify_mod, alpha, alpha_feature(integration=[]))
    assert r["level"] == 2 and any(x.startswith("failing") for x in r["reasons"])


@pytest.mark.parametrize("name,body,why", [
    ("assert_true", "from sidecar.alpha import compute\n\n\ndef test_x():\n    compute(1, 2)\n    assert True\n", "no real assertions"),
    ("self_compare", "from sidecar.alpha import compute\n\n\ndef test_x():\n    assert compute(1, 2) == compute(1, 2)\n", "no real assertions"),
    ("constant_one", "from sidecar.alpha import compute\n\n\ndef test_x():\n    compute(1, 2)\n    assert 1\n", "no real assertions"),
    ("no_tests", "from sidecar.alpha import compute\n\nVALUE = compute(1, 2)\n", "no test functions"),
    ("unrelated", "def test_math():\n    assert 1 + 1 == 2\n", "mentions what the feature implements"),
])
def test_vacuous_tests_do_not_count(verify_mod, alpha, name, body, why):
    put(alpha, "tests/unit/test_alpha.py", body)
    r = ev(verify_mod, alpha, alpha_feature(integration=[]))
    assert r["level"] == 2 and "VACUOUS_TESTS" in r["flags"], (name, r["reasons"])
    assert any(why in x for x in r["reasons"]), r["reasons"]


def test_skipped_tests_are_not_a_pass(verify_mod, alpha):
    put(alpha, "tests/unit/test_alpha.py",
        "import pytest\nfrom sidecar.alpha import compute\n\n\n@pytest.mark.skip(reason='later')\ndef test_x():\n    assert compute(1, 2) == 3\n")
    r = ev(verify_mod, alpha, alpha_feature(integration=[]))
    assert r["level"] == 2 and any(x.startswith("failing") for x in r["reasons"])


def test_a_command_that_names_no_test_file_does_not_count(verify_mod, alpha):
    f = alpha_feature(integration=[], unit=[f"{QPY} -m pytest -q"])
    r = ev(verify_mod, alpha, f)
    assert r["level"] == 2 and "VACUOUS_TESTS" in r["flags"] and any("names no existing test" in x for x in r["reasons"])


def test_echo_is_not_a_test(verify_mod, alpha):
    r = ev(verify_mod, alpha, alpha_feature(integration=[], unit=["echo ok"]))
    assert r["level"] == 2 and any("not a recognised test runner" in x for x in r["reasons"])


def test_custom_script_counts_only_with_an_expected_success_line(verify_mod, alpha):
    put(alpha, "scripts/e2e.py", "import pathlib\nimport sys\n\nsys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))\n"
                                 "from sidecar.alpha import compute\n\nassert compute(1, 2) == 3\nprint('E2E OK')\n")
    cmd = f"{QPY} scripts/e2e.py"
    assert ev(verify_mod, alpha, alpha_feature(integration=[], unit=[{"cmd": cmd, "expect_regex": "E2E OK"}]))["level"] == 3
    assert ev(verify_mod, alpha, alpha_feature(integration=[], unit=[{"cmd": cmd, "expect_regex": "NEVER PRINTED"}]))["level"] == 2
    assert ev(verify_mod, alpha, alpha_feature(integration=[], unit=[cmd]))["level"] == 2


def test_unit_run_directory_form_is_analysed_too(verify_mod, alpha):
    put(alpha, "tests/unit/test_alpha.py", "from sidecar.alpha import compute\n\n\ndef test_x():\n    compute(1, 2)\n    assert True\n")
    r = ev(verify_mod, alpha, alpha_feature(integration=[], unit=[pytest_cmd("tests/unit")]))
    assert r["level"] == 2 and "VACUOUS_TESTS" in r["flags"]


def test_integration_reaches_level_4_and_must_differ_from_the_unit_command(verify_mod, alpha):
    r = ev(verify_mod, alpha, alpha_feature())
    assert (r["label"], r["pct"]) == ("INTEGRATED", 85), r["reasons"]
    same = alpha_feature(integration=[pytest_cmd("tests/unit/test_alpha.py")])
    r2 = ev(verify_mod, alpha, same)
    assert r2["level"] == 3 and any("identical" in x for x in r2["reasons"])


def test_failing_integration_stays_at_level_3(verify_mod, alpha):
    put(alpha, "tests/integration/test_alpha_flow.py", "from sidecar.alpha import compute\n\n\ndef test_flow():\n    assert compute(1, 1) == 3\n")
    r = ev(verify_mod, alpha, alpha_feature())
    assert r["level"] == 3 and any("integration failing" in x for x in r["reasons"])


def test_guard_evidence(verify_mod, repo):
    f = feat(files=["sidecar/gamma.py"], symbols=[{"file": "sidecar/gamma.py", "name": "gamma"}], guard_ids=["forbidden-path", "py-import-core"])
    ids = verify_mod.known_guard_ids({"forbidden_text": [{"id": "py-import-core"}], "stale_paths": []})
    clean = {"findings": []}
    assert ev(verify_mod, repo, f, guard_report=clean, ids=ids)["level"] == 3
    dirty = {"findings": [{"check": "py-import-core", "severity": "FAIL", "message": "import core", "where": "x"}]}
    r = ev(verify_mod, repo, f, guard_report=dirty, ids=ids)
    assert r["level"] == 2 and any("py-import-core" in x for x in r["reasons"])
    assert ev(verify_mod, repo, f, guard_report=None, ids=ids)["level"] == 2


def test_a_mistyped_guard_id_cannot_pass_vacuously(verify_mod, repo):
    f = feat(files=["sidecar/gamma.py"], symbols=[{"file": "sidecar/gamma.py", "name": "gamma"}], guard_ids=["forbiden-path"])
    r = ev(verify_mod, repo, f, guard_report={"findings": []}, ids=verify_mod.known_guard_ids({}))
    assert r["level"] == 2 and any("unknown guard check id" in str(x) for x in r["checks"]["guard"])


def test_every_finding_id_emitted_by_the_guard_is_in_the_shared_list(common):
    src = (TOOLS / "guard" / "check.py").read_text(encoding="utf-8")
    emitted = set(re.findall(r'Finding\("(?:FAIL|WARN|INFO)",\s*"([a-z][a-z\-]*)"', src)) | {"human-review"}
    emitted -= {"layout"}  # dynamic names (crash reports) are never cited by evidence
    missing = emitted - common.GUARD_CHECK_IDS
    assert not missing, f"add to locus_common.GUARD_CHECK_IDS: {sorted(missing)}"


# ---- proof, signoff, staleness
def record(repo, fid, *cmd) -> subprocess.CompletedProcess:
    args = [fid]
    if cmd:
        args += ["--", *cmd]
    return tracker(repo, "record_proof.py", *args)


def proof_id(repo, fid) -> str:
    return json.loads((repo / "tracker" / "proof" / f"{fid}.json").read_text(encoding="utf-8"))["files_sha256"][:12]


def put_feature_in_yaml(repo, f):
    p = repo / "tracker" / "features.yaml"
    data = yaml.safe_load(p.read_text(encoding="utf-8"))
    data["features"].append(f)
    p.write_text(yaml.safe_dump(data, sort_keys=False, allow_unicode=True), encoding="utf-8")


def test_proof_and_signoff_reach_level_5_and_go_stale_when_code_changes(verify_mod, alpha):
    f = alpha_feature()
    put_feature_in_yaml(alpha, f)
    assert ev(verify_mod, alpha, f)["level"] == 4
    assert any("no recorded proof" in x for x in ev(verify_mod, alpha, f)["reasons"])

    p = record(alpha, "F-9001", PY, "-c", "print('e2e ok')")
    assert p.returncode == 0 and 'proof: "' in p.stdout
    r = ev(verify_mod, alpha, f)
    assert r["level"] == 4 and any("waiting for human signoff" in x for x in r["reasons"])

    pid = proof_id(alpha, "F-9001")
    sign = {"F-9001": {"by": "Shivam", "date": "2026-10-05", "proof": pid, "note": "watched it"}}
    r = ev(verify_mod, alpha, f, signoff=sign)
    assert (r["label"], r["pct"]) == ("PROVEN", 100)

    put(alpha, "sidecar/alpha.py", ALPHA_PY + "\n# touched\n")
    r = ev(verify_mod, alpha, f, signoff=sign)
    assert r["level"] == 4 and "PROOF_STALE" in r["flags"]

    record(alpha, "F-9001", PY, "-c", "print('e2e ok again')")  # an agent re-records the proof...
    r = ev(verify_mod, alpha, f, signoff=sign)  # ...but the old signoff names the old proof
    assert r["level"] == 4 and "SIGNOFF_STALE" in r["flags"]
    sign["F-9001"]["proof"] = proof_id(alpha, "F-9001")
    assert ev(verify_mod, alpha, f, signoff=sign)["level"] == 5


@pytest.mark.parametrize("entry", [{}, {"by": "S"}, {"by": "S", "date": "2026-01-01"}, "yes"])
def test_malformed_signoff_does_not_count(verify_mod, alpha, entry):
    f = alpha_feature()
    put_feature_in_yaml(alpha, f)
    record(alpha, "F-9001", PY, "-c", "print(1)")
    r = ev(verify_mod, alpha, f, signoff={"F-9001": entry})
    assert r["level"] == 4 and any("must have by:, date: and proof:" in x for x in r["reasons"])


def test_failed_proof_run_is_recorded_and_does_not_count(verify_mod, alpha):
    f = alpha_feature()
    put_feature_in_yaml(alpha, f)
    p = record(alpha, "F-9001", PY, "-c", "import sys; sys.exit(3)")
    assert p.returncode == 1 and "signoff.yaml" not in p.stdout
    r = ev(verify_mod, alpha, f, signoff={"F-9001": {"by": "S", "date": "d", "proof": proof_id(alpha, "F-9001")}})
    assert r["level"] == 4 and any("proof run failed" in x for x in r["reasons"])


def test_manual_proof_needs_note_and_artifact(alpha):
    put_feature_in_yaml(alpha, alpha_feature())
    assert record(alpha, "F-9001").returncode == 2
    assert tracker(alpha, "record_proof.py", "F-9001", "--note", "saw it", "--artifact", "nope.png").returncode == 2
    put(alpha, "evidence/shot.txt", "screenshot stand-in\n")
    p = tracker(alpha, "record_proof.py", "F-9001", "--note", "saw it", "--artifact", "evidence/shot.txt")
    assert p.returncode == 0
    d = json.loads((alpha / "tracker" / "proof" / "F-9001.json").read_text(encoding="utf-8"))
    assert d["manual"] is True and d["artifact_sha256"]


def test_record_proof_rejects_unknown_feature_or_feature_without_files(alpha):
    assert record(alpha, "F-0000", PY, "-c", "print(1)").returncode == 2
    put_feature_in_yaml(alpha, feat(fid="F-9100"))
    assert record(alpha, "F-9100", PY, "-c", "print(1)").returncode == 2


def test_manual_features_skip_tests_but_still_need_proof_and_signoff(verify_mod, repo):
    f = feat(fid="F-9002", files=["sidecar/gamma.py"], symbols=[{"file": "sidecar/gamma.py", "name": "gamma"}], manual=True)
    put_feature_in_yaml(repo, f)
    r = ev(verify_mod, repo, f)
    assert r["level"] == 2 and any("no recorded proof" in x for x in r["reasons"])
    record(repo, "F-9002", PY, "-c", "print('checked by hand')")
    sign = {"F-9002": {"by": "S", "date": "2026-01-01", "proof": proof_id(repo, "F-9002")}}
    assert ev(verify_mod, repo, f, signoff=sign)["level"] == 5


# ---- inherited, overclaim, caps, spec drift
def test_inherited_pi_feature(verify_mod, repo):
    f = feat(files=["packages/pi-mono/index.ts"], inherited=True)
    r = ev(verify_mod, repo, f, upstream_ok=True)
    assert (r["label"], r["pct"]) == ("INHERITED_OK", 100)
    r2 = ev(verify_mod, repo, f, upstream_ok=False)
    assert r2["level"] == 1 and any("upstream pin" in x for x in r2["reasons"])


def test_inherited_flag_on_locus_code_is_ignored(verify_mod, repo):
    f = feat(files=["sidecar/gamma.py"], inherited=True)
    r = ev(verify_mod, repo, f, upstream_ok=True)
    assert r["label"] == "FILES_PRESENT" and "INHERITED_INVALID" in r["flags"]
    mixed = feat(files=["packages/pi-mono/index.ts", "sidecar/gamma.py"], inherited=True)
    assert "INHERITED_INVALID" in ev(verify_mod, repo, mixed, upstream_ok=True)["flags"]


def test_overclaim_flag(verify_mod, repo):
    r = ev(verify_mod, repo, feat(files=["sidecar/gamma.py"], claimed={"pct": 90}))
    assert r["gap"] == 75.0 and "OVERCLAIM" in r["flags"]
    assert "OVERCLAIM" not in ev(verify_mod, repo, feat(files=["sidecar/gamma.py"], claimed={"pct": 30}))["flags"]


def test_human_cap_lowers_a_level(verify_mod, alpha):
    r = ev(verify_mod, alpha, alpha_feature(), overrides={"F-9001": {"cap_level": 1, "by": "S", "reason": "x"}})
    assert r["level"] == 1 and "CAPPED_BY_HUMAN" in r["flags"]


def test_spec_drift_flags(verify_mod, repo):
    f = feat(files=["sidecar/gamma.py"], spec_sha={"§1.1: Alpha tool": "old"})
    assert "SPEC_CHANGED" in ev(verify_mod, repo, f, spec_sha={"1.1 alpha tool": "new"})["flags"]
    assert "SPEC_REF_DANGLING" in ev(verify_mod, repo, f, spec_sha={})["flags"]
    assert not {"SPEC_CHANGED", "SPEC_REF_DANGLING"} & set(ev(verify_mod, repo, f, spec_sha={"1.1 alpha tool": "old"})["flags"])


def test_js_test_analysis(verify_mod, repo):
    good = "import { alpha } from '../src/tools/alpha';\nit('works', () => {\n  expect(alpha()).toBe(1);\n});\n"
    triv = "import { alpha } from '../src/tools/alpha';\nit('works', () => {\n  alpha();\n  expect(true).toBe(true);\n});\n"
    skipped = "import { alpha } from '../src/tools/alpha';\nit.skip('works', () => {\n  expect(alpha()).toBe(1);\n});\n"
    comment = "import { alpha } from '../src/tools/alpha';\n// expect(alpha()).toBe(1);\nit('works', () => { alpha(); });\n"
    cmd = {"cmd": "npx vitest run locus-extension/test/alpha.test.ts"}
    for name, body, vacuous in (("good", good, False), ("trivial", triv, True), ("skipped", skipped, True), ("comment", comment, True)):
        put(repo, "locus-extension/test/alpha.test.ts", body)
        runner, vac, why = verify_mod.judge_test_command(repo, cmd, {"alpha"})
        assert runner == "vitest" and vac is vacuous, (name, why)


@pytest.mark.parametrize("cmd,runner", [
    ("python -m pytest tests/x.py", "pytest"), ("pytest -q", "pytest"), ("python -m unittest tests.test_x", "unittest"),
    ("npx vitest run", "vitest"), ("npx jest --ci", "jest"), ("node --test tests/", "node-test"), ("echo hi", None), ("npm test", None),
])
def test_runner_detection(verify_mod, cmd, runner):
    assert verify_mod.detect_runner(cmd) == runner


# ---- verify.py as a program, then the workbook
def test_verify_cli_writes_status_markdown_and_history(repo):
    p = tracker(repo, "verify.py", "--static")
    assert p.returncode == 0, p.stdout + p.stderr
    st = json.loads((repo / "tracker" / "status.json").read_text(encoding="utf-8"))
    levels = {k: v["label"] for k, v in st["features"].items()}
    assert levels == {"F-0001": "FILES_PRESENT", "F-0002": "INHERITED_OK", "F-0003": "IMPLEMENTED"}
    assert st["features"]["F-0001"]["flags"].count("OVERCLAIM") == 1
    assert (repo / "tracker" / "STATUS.md").is_file()
    assert len((repo / "tracker" / "history.jsonl").read_text(encoding="utf-8").splitlines()) == 1
    assert st["guard"]["verdict"] == "PASS"


def test_modified_pi_code_removes_inherited_status(repo):
    put(repo, "packages/pi-mono/index.ts", "export const upstream = 99;\n")
    tracker(repo, "verify.py", "--static")
    st = json.loads((repo / "tracker" / "status.json").read_text(encoding="utf-8"))
    assert st["features"]["F-0002"]["label"] == "FILES_PRESENT"


def test_regression_is_reported_and_can_fail_ci(repo):
    tracker(repo, "verify.py", "--static")
    (repo / "sidecar" / "gamma.py").unlink()
    p = tracker(repo, "verify.py", "--static", "--fail-on-regression")
    assert p.returncode == 1 and "REGRESSION" in p.stdout and "F-0003" in p.stdout


def test_static_and_full_runs_are_never_compared(repo):
    hist = repo / "tracker" / "history.jsonl"
    hist.write_text(json.dumps({"ts": "t", "commit": "c", "dirty": False, "static": False, "mean": 100, "levels": {"F-0003": 100}}) + "\n",
                    encoding="utf-8")
    p = tracker(repo, "verify.py", "--static", "--fail-on-regression")
    assert p.returncode == 0 and "REGRESSION" not in p.stdout  # the only history entry is from a full run
    (repo / "sidecar" / "gamma.py").unlink()
    p2 = tracker(repo, "verify.py", "--static", "--fail-on-regression")
    assert p2.returncode == 1 and "F-0003" in p2.stdout


def test_fail_on_guard(repo):
    put(repo, "core/x.py", "x = 1\n")
    assert tracker(repo, "verify.py", "--static", "--fail-on-guard").returncode == 1


def test_only_flag_does_not_overwrite_status_files(repo):
    p = tracker(repo, "verify.py", "--static", "--only", "F-0003")
    assert p.returncode == 0 and '"F-0003"' in p.stdout
    assert not (repo / "tracker" / "status.json").exists()


def test_show_prints_feature_status_and_spec_translation(repo):
    tracker(repo, "verify.py", "--static")
    p = tracker(repo, "show.py", "F-0001", "--spec")
    assert p.returncode == 0
    assert "COMPUTED STATUS" in p.stdout and "Lives in" in p.stdout and "FILES_PRESENT" in p.stdout
    f = tracker(repo, "show.py", "--find", "gamma")
    assert "F-0003" in f.stdout
    assert tracker(repo, "show.py", "F-9999").returncode == 2


def test_export_builds_a_generated_workbook_that_the_guard_accepts(repo):
    import openpyxl  # noqa: PLC0415
    tracker(repo, "verify.py", "--static")
    p = tracker(repo, "export_xlsx.py")
    assert p.returncode == 0, p.stdout + p.stderr
    xlsx = repo / "spec" / "Tracker.xlsx"
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames == ["Feature Tracker", "Summary", "Legend", "Spec Coverage", "_meta"]
    assert wb["_meta"].sheet_state == "hidden"
    ws = wb["Feature Tracker"]
    assert ws.max_row == 4 and ws["A2"].value == "F-0001" and ws["I3"].value == "INHERITED_OK"
    assert guard(repo, "--full").of("tracker-xlsx") == []
    wb.save(xlsx)  # opening and saving without edits (what Excel does) must not trip the tamper check
    assert guard(repo, "--full").of("tracker-xlsx", "FAIL") == []


def test_hand_editing_the_workbook_is_detected(repo):
    import openpyxl  # noqa: PLC0415
    tracker(repo, "verify.py", "--static")
    tracker(repo, "export_xlsx.py")
    xlsx = repo / "spec" / "Tracker.xlsx"
    wb = openpyxl.load_workbook(xlsx)
    wb["Feature Tracker"]["J2"] = 100  # "make F-0001 look finished"
    wb.save(xlsx)
    r = guard(repo, "--full")
    assert r.rc == 1 and any("hand-edited" in f["message"] for f in r.of("tracker-xlsx", "FAIL"))


def test_old_hand_maintained_workbook_is_replaced_and_backed_up(repo):
    import openpyxl  # noqa: PLC0415
    xlsx = repo / "spec" / "Tracker.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Feature Tracker"
    wb.active.append(["§ID", "Completion %"])
    wb.active.append(["§1.1", 100])
    wb.save(xlsx)
    assert guard(repo, "--full").of("tracker-xlsx", "WARN")
    tracker(repo, "verify.py", "--static")
    p = tracker(repo, "export_xlsx.py")
    assert "backed up" in p.stdout
    assert list((repo / "tracker" / "legacy").glob("replaced-*Tracker.xlsx"))
    assert guard(repo, "--full").of("tracker-xlsx") == []


@pytest.mark.skipif(shutil.which("soffice") is None and shutil.which("libreoffice") is None, reason="LibreOffice not installed")
def test_summary_formulas_calculate_without_errors(repo, tmp_path):
    import openpyxl  # noqa: PLC0415
    tracker(repo, "verify.py", "--static")
    tracker(repo, "export_xlsx.py")
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    out = tmp_path / "calc"
    subprocess.run([exe, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(repo / "spec" / "Tracker.xlsx")],
                   capture_output=True, timeout=180, env={**os.environ, "HOME": str(tmp_path)})
    calc = openpyxl.load_workbook(out / "Tracker.xlsx", data_only=True)["Summary"]
    values = {calc.cell(r, 1).value: calc.cell(r, 2).value for r in range(1, calc.max_row + 1) if calc.cell(r, 1).value}
    assert values["Features tracked"] == 3
    assert values["Verified average %"] == pytest.approx(51.7, abs=0.05)
    assert values["Level INHERITED_OK"] == 1 and values["Level FILES_PRESENT"] == 1 and values["Level IMPLEMENTED"] == 1
    assert not [v for v in values.values() if isinstance(v, str) and v.startswith("#")]


def test_average_formulas_survive_an_empty_claimed_column(repo):
    import openpyxl  # noqa: PLC0415
    edit_features(repo, lambda fs: [f.pop("claimed") for f in fs])
    tracker(repo, "verify.py", "--static")
    tracker(repo, "export_xlsx.py")
    sm = openpyxl.load_workbook(repo / "spec" / "Tracker.xlsx")["Summary"]
    formulas = [sm.cell(r, 2).value for r in range(1, sm.max_row + 1) if isinstance(sm.cell(r, 2).value, str) and "AVERAGE" in sm.cell(r, 2).value]
    assert formulas and all(f.startswith("=IFERROR(") for f in formulas)


# ================================================================== 12. importing the old workbook
LEGACY_HEADERS = ["§ID", "Feature Name", "Category", "Verdict", "Source Codebase / Stealing From", "Provider Interface", "Pi Agent Status",
                  "Sidecar Status", "Completion %", "Remaining Work", "Key Files", "Notes / Issues"]


def make_legacy(path: Path) -> None:
    import openpyxl  # noqa: PLC0415
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feature Tracker"
    ws.append(LEGACY_HEADERS)
    ws.append(["§1.1", "Alpha tool", "Foundation", "LOCUS-NEW", "Locus", "—", "✅", "✅", 100, "None", "locus-extension/src/tools/alpha.ts", ""])
    ws.append(["§1.2", "Beta engine", "Foundation", "KEEP-PI", "Pi", "—", "✅ INHERITED", "N/A", 100, "None", "packages/pi-mono/index.ts", ""])
    ws.append(["ARCH-3", "Python package is sidecar", "Infrastructure", "LOCUS-NEW", "Locus", "—", "n/a", "✅", 60, "x", "sidecar/__init__.py", ""])
    ws.append(["§2.1", "Totally unrelated words here", "Sidecar", "LOCUS-NEW", "Locus", "—", "n/a", "✅", 80, "x", "sidecar/gamma.py", ""])
    wb.create_sheet("Summary")
    wb.save(path)


def test_import_legacy_creates_evidence_only_features(repo, tmp_path):
    make_legacy(tmp_path / "old.xlsx")
    (repo / "tracker" / "features.yaml").unlink()
    spec = repo / "spec" / "SPEC.md"
    spec.write_text(spec.read_text(encoding="utf-8") + "\n## Feature §3.1: Delta\n**Verdict**: LOCUS-NEW\n", encoding="utf-8")
    p = tracker(repo, "import_legacy.py", "--xlsx", str(tmp_path / "old.xlsx"))
    assert p.returncode == 0, p.stdout + p.stderr
    data = yaml.safe_load((repo / "tracker" / "features.yaml").read_text(encoding="utf-8"))
    feats = {f["title"]: f for f in data["features"]}
    alpha_f = feats["Alpha tool"]
    assert alpha_f["claimed"]["pct"] == 100 and alpha_f["evidence"]["files"] == ["locus-extension/src/tools/alpha.ts"]
    assert not ({"status", "pct", "percent", "level"} & set(alpha_f))
    assert feats["Beta engine"]["inherited"] is True
    assert feats["Totally unrelated words here"]["mapping"] == "low-confidence"
    arch = feats["Python package is sidecar"]
    assert "sidecar/__init__.py" in arch["evidence"]["files"] and "py-import-core" in arch["evidence"]["guard"]
    assert any(f["title"] == "Delta" and f["origin"].startswith("spec-gap") for f in data["features"])  # no legacy row covered §3.1
    assert (repo / "tracker" / "legacy" / "legacy_claims.json").is_file()


def test_import_legacy_refuses_to_overwrite_evidence(repo, tmp_path):
    make_legacy(tmp_path / "old.xlsx")
    p = tracker(repo, "import_legacy.py", "--xlsx", str(tmp_path / "old.xlsx"))
    assert p.returncode == 2 and "Refusing" in p.stdout
    assert tracker(repo, "import_legacy.py", "--xlsx", str(tmp_path / "old.xlsx"), "--force").returncode == 0


# ================================================================== 13. shared helpers
@pytest.mark.parametrize("text,ids", [("§3.7 + 11.5", ["3.7", "11.5"]), ("17.1-17.4", ["17.1", "17.2", "17.3", "17.4"]), ("§4.2–4.3", ["4.2", "4.3"]),
                                       ("1.5a", ["1.5"]), ("ARCH-1", [])])
def test_expand_ids(common, text, ids):
    assert common.expand_ids(text) == ids


def test_parse_spec_sections_hashes_and_targets(common, repo):
    spec = common.parse_spec(repo / "spec" / "SPEC.md")
    heads = spec["headings"]
    assert [h["key"] for h in heads] == ["1.1 alpha tool", "1.2 beta engine", "2.1 gamma sidecar"]
    assert heads[0]["verdict"] == "LOCUS-NEW" and heads[0]["targets"] == ["locus-extension/src/tools/alpha.ts"]
    before = heads[1]["sha"]
    (repo / "spec" / "SPEC.md").write_text((repo / "spec" / "SPEC.md").read_text(encoding="utf-8").replace("KEEP-PI", "AUGMENT"), encoding="utf-8")
    after = common.parse_spec(repo / "spec" / "SPEC.md")["headings"]
    assert after[1]["sha"] != before and after[0]["sha"] == heads[0]["sha"]


def test_translate_target_uses_the_stale_path_rules(common):
    rules = common.load_toml(TOOLS / "guard" / "guard.toml")["stale_paths"]
    t = common.translate_target
    assert t("locus-os/src/core/tools/x.ts", rules) == "locus-extension/src/tools/x.ts"
    assert t("locus/core/x.ts", rules) == "locus-extension/src/x.ts"
    assert t("locus/extensions/y.ts", rules) == "locus-extension/src/y.ts"
    assert t("src/extensions/z.ts", rules) == "locus-extension/src/z.ts"
    assert t("locus-os/sidecar/x.py", rules) == "sidecar/x.py"
    assert t("sidecar/x.py", rules) == "sidecar/x.py"


def test_grid_hash_ignores_int_float_noise_and_blank_rows(common):
    import openpyxl  # noqa: PLC0415
    a, b = openpyxl.Workbook().active, openpyxl.Workbook().active
    for ws, v in ((a, 15), (b, 15.0)):
        ws.append(["h"] * 16)
        ws.append(["F-1"] + [None] * 8 + [v] + [None] * 6)
    b.append([None] * 16)
    assert common.grid_hash(a) == common.grid_hash(b)
    b["J2"] = 16
    assert common.grid_hash(a) != common.grid_hash(b)


def test_sha256_paths_changes_with_content_and_ignores_pycache(common, tmp_path):
    (tmp_path / "d").mkdir()
    (tmp_path / "d" / "a.py").write_text("1", encoding="utf-8")
    h1 = common.sha256_paths(tmp_path, ["d"])
    (tmp_path / "d" / "__pycache__").mkdir()
    (tmp_path / "d" / "__pycache__" / "a.pyc").write_text("x", encoding="utf-8")
    assert common.sha256_paths(tmp_path, ["d"]) == h1
    (tmp_path / "d" / "a.py").write_text("2", encoding="utf-8")
    assert common.sha256_paths(tmp_path, ["d"]) != h1


# ================================================================== 14. the guard.toml rules themselves
def test_guard_toml_parses_and_every_regex_compiles(common):
    cfg = common.load_toml(TOOLS / "guard" / "guard.toml")
    pats = [r["pattern"] for r in cfg["forbidden_text"]] + [r["pattern"] for r in cfg["stale_paths"]] + [p["pattern"] for p in cfg["secrets"]["patterns"]]
    pats += [p["pattern"] for p in cfg["spec_lint"]["required_principles"]]
    for p in pats:
        re.compile(p)
    ids = [r["id"] for r in cfg["forbidden_text"]] + [r["id"] for r in cfg["stale_paths"]]
    assert len(ids) == len(set(ids)), "rule ids must be unique"
    assert not set(ids) & set(__import__("locus_common").GUARD_CHECK_IDS)


@pytest.mark.parametrize("text", ["see locus-os/src/core/x.ts", "locus/core/y.ts", "locus/extensions/z.ts", "src/extensions/w.ts"])
def test_stale_path_patterns_match_spec_style_paths(common, text):
    rules = common.load_toml(TOOLS / "guard" / "guard.toml")["stale_paths"]
    assert any(re.search(r["pattern"], text) for r in rules)


@pytest.mark.parametrize("text", ["locus-extension/src/core/x.ts", "my-locus/core/x.ts", "packages/locus/core/x.ts", "locus-extension/src/tools/x.ts"])
def test_stale_path_patterns_ignore_the_canonical_layout(common, text):
    rules = common.load_toml(TOOLS / "guard" / "guard.toml")["stale_paths"]
    assert not any(re.search(r["pattern"], text) for r in rules)


def test_vault_token_regex(check_mod):
    rx = check_mod.VAULT_TOKEN
    assert [m.group(1) for m in rx.finditer('x = "05-Logs/a" + "06-Resources"')] == ["05-Logs", "06-Resources"]
    assert not rx.findall("2024-01-15 and 2026-10-04T05:06 and v01-alpha and file.01-x")


# ================================================================== 15. the bundle's own rule files obey the guard's limits
def test_bundle_rule_files_obey_the_size_and_import_rules(tmp_path):
    if not (BUNDLE / "AGENTS.md").is_file():
        pytest.skip("AGENTS.md not installed yet")
    out = tmp_path / "r.json"
    p = subprocess.run([PY, str(TOOLS / "guard" / "check.py"), "--root", str(BUNDLE), "--only", "agent-files", "--json", str(out)],
                       capture_output=True, text=True, encoding="utf-8", env=base_env())
    rep = json.loads(out.read_text(encoding="utf-8"))
    assert [f for f in rep["findings"] if f["severity"] == "FAIL"] == [], p.stdout


def test_documentation_only_names_tools_and_segments_that_exist():
    docs = [p for p in [BUNDLE / "AGENTS.md", BUNDLE / "CLAUDE.md", BUNDLE / "README.md", BUNDLE / "RUNBOOK.md"] if p.is_file()]
    docs += sorted((BUNDLE / "docs" / "agent").rglob("*.md")) if (BUNDLE / "docs" / "agent").is_dir() else []
    if not docs:
        pytest.skip("no documentation installed yet")
    problems = []
    for d in docs:
        text = d.read_text(encoding="utf-8")
        for m in re.finditer(r"(?<![\w/.-])(tools/[\w\-/]+\.py)\b", text):
            if not (BUNDLE / m.group(1)).is_file():
                problems.append(f"{d.relative_to(BUNDLE)} mentions {m.group(1)} which does not exist")
        for m in re.finditer(r"(?<![\w/.-])(docs/agent/(?:prompts/)?[\w\-]+\.md)\b", text):
            if not (BUNDLE / m.group(1)).is_file():
                problems.append(f"{d.relative_to(BUNDLE)} mentions {m.group(1)} which does not exist")
        # short names (`08-task-protocol.md`, `P1-implement-feature.md`) must exist as a segment or a prompt
        for m in re.finditer(r"(?<![\w/.\\-])((?:\d\d|P\d)-[a-z0-9\-]+\.md)\b", text):
            name = m.group(1)
            if not ((BUNDLE / "docs" / "agent" / name).is_file() or (BUNDLE / "docs" / "agent" / "prompts" / name).is_file()):
                problems.append(f"{d.relative_to(BUNDLE)} mentions {name} which does not exist")
    assert not problems, "\n".join(problems)


# ================================================================== 16. the Agile board (tools/tracker/backlog.py)
AGILE_ERRATA = """# errata
## A. Index
| ID | Status | Topic | Blocks |
|---|---|---|---|
| D-1 | OPEN | A formula | x |
| D-2 | DECIDED | Vault | - |
| D-10 | OPEN (low urgency) | Config location | y |
"""

AGILE_BACKLOG = """schema: 1
config:
  first_sprint_start: 2026-10-16
  capacity_points: 4
epics:
  - {id: E1, title: One, goal: "Goal one"}
stories:
  - {id: US-01, epic: E1, feature: F-0001, slug: alpha-one, title: "Alpha story", points: 2, target_level: 3, accept: ["Alpha works"], demo: "Run alpha"}
  - {id: US-02, epic: E1, feature: F-0003, slug: gamma-two, title: "Gamma story", points: 1, target_level: 3, accept: ["Gamma works"], demo: "Run gamma"}
  - {id: US-03, epic: E1, feature: F-0003, slug: gamma-three, title: "Needs a decision", points: 2, target_level: 3, requires_decisions: [D-10], accept: ["x"], demo: "y"}
  - {id: US-04, epic: E1, feature: F-0001, slug: alpha-four, title: "Needs US-01", points: 2, target_level: 3, depends_on: [US-01], accept: ["z"], demo: "w"}
  - {id: US-05, epic: E1, feature: F-0001, slug: coarse, title: "Coarse", points: 3, target_level: 3}
"""

AGILE_HANDOFF = """# US-01 handoff - claude sonnet, 2026-10-19
## Plan (written before coding)
the plan
## Files changed (every path, one line each)
sidecar/alpha.py
## Spec sentences implemented (quote + spec line) / NOT implemented and why
quoted
## Config keys and defaults added
none
## Commands run, with real output (guard, verify line, tests)
ok
## Open questions / errata rows touched
none
## Discard table (path | why | replaced by | imported by)
none
## Security (anything found, even if outside this story)
none
## Not done (be specific)
nothing
"""
AGILE_VERIFY_OK = "VERDICT: ACCEPT\nVerifier: Gemini 3 Pro\n## Findings\nnone\n"


def agile(repo: Path, level: int = 3) -> Path:
    shutil.copytree(BUNDLE / "docs" / "agent" / "prompts", repo / "docs" / "agent" / "prompts")
    put(repo, "docs/agent/09-spec-errata.md", AGILE_ERRATA)
    put(repo, "tracker/backlog.yaml", AGILE_BACKLOG)
    put(repo, "tracker/status.json", json.dumps({"features": {"F-0001": {"level": level}, "F-0003": {"level": level}}}))
    commit_all(repo, "agile fixtures")
    return repo


def bl(repo: Path, *args: str, today: str = "2026-10-16", owner: bool = False) -> subprocess.CompletedProcess:
    return tracker(repo, "backlog.py", *args, "--today", today, env={"LOCUS_ALLOW_GUARD_EDIT": "1"} if owner else None)


def story_states(repo: Path, today: str = "2026-10-19") -> dict:
    p = bl(repo, "--json", today=today)
    assert p.returncode == 0, p.stdout + p.stderr
    return json.loads(p.stdout)["states"]


def plan_sprint(repo: Path) -> None:
    p = bl(repo, "--plan", "--write", owner=True)
    assert p.returncode == 0, p.stdout + p.stderr
    commit_all(repo, "plan sprint")


def merged_story(repo: Path, sid: str = "US-01", slug: str = "alpha-one") -> None:
    p = bl(repo, "--start", sid)
    assert p.returncode == 0, p.stdout + p.stderr
    put(repo, f"tracker/handoff/{sid}.md", AGILE_HANDOFF.replace("US-01", sid))
    put(repo, f"tracker/handoff/{sid}-verify.md", AGILE_VERIFY_OK)
    commit_all(repo, f"feat({sid}): work")
    git(repo, "switch", "-q", "main")
    git(repo, "merge", "--ff-only", f"task/{sid}-{slug}")
    git(repo, "branch", "-d", f"task/{sid}-{slug}")


def test_backlog_check_passes_and_warns_about_top_stories_that_are_not_ready(repo):
    agile(repo)
    p = bl(repo, "--check")
    assert p.returncode == 0, p.stdout
    assert "0 ERROR" in p.stdout and "US-03 is not Ready" in p.stdout


def test_backlog_check_rejects_structural_mistakes(repo):
    agile(repo)
    bad = (AGILE_BACKLOG.replace("feature: F-0003, slug: gamma-two", "feature: F-9999, slug: gamma-two")
           .replace('points: 2, target_level: 3, accept: ["Alpha works"]', 'points: 5, target_level: 3, accept: ["Alpha works"]')
           .replace("depends_on: [US-01]", "depends_on: [US-06]")
           .replace("requires_decisions: [D-10]", "requires_decisions: [D-99]"))
    bad += '  - {id: US-01, epic: E1, feature: F-0001, slug: dup, title: "Dup", points: 1, target_level: 3}\n'
    put(repo, "tracker/backlog.yaml", bad)
    p = bl(repo, "--check")
    assert p.returncode == 1
    for needle in ("feature F-9999", "points must be 1 to 3", "depends_on US-06 does not exist", "D-99 is not a row", "duplicate id"):
        assert needle in p.stdout, (needle, p.stdout)


def test_a_story_must_come_after_what_it_depends_on(repo):
    agile(repo)
    put(repo, "tracker/backlog.yaml", AGILE_BACKLOG.replace("depends_on: [US-01]", "depends_on: [US-05]"))
    p = bl(repo, "--check")
    assert p.returncode == 1 and "not ABOVE it" in p.stdout


def test_plan_takes_ready_stories_in_priority_order_and_skips_the_rest(repo):
    agile(repo)
    p = bl(repo, "--plan")
    assert p.returncode == 0, p.stdout
    assert "take    US-01" in p.stdout and "take    US-02" in p.stdout and "total 3/4" in p.stdout
    assert "skipped US-03" in p.stdout and "D-10 is OPEN" in p.stdout
    assert "skipped US-04" in p.stdout and "waits for US-01" in p.stdout
    assert "skipped US-05" in p.stdout
    assert not (repo / "tracker" / "sprints.jsonl").exists()


def test_planning_stops_at_the_first_story_that_does_not_fit(repo):
    agile(repo)
    put(repo, "tracker/backlog.yaml", AGILE_BACKLOG.replace("points: 1, target_level: 3, accept: [\"Gamma works\"]", "points: 3, target_level: 3, accept: [\"Gamma works\"]"))
    p = bl(repo, "--plan")
    assert "take    US-01" in p.stdout and "stop    US-02" in p.stdout and "take    US-02" not in p.stdout


def test_recording_a_plan_needs_the_owner_a_slot_day_and_the_planning_date(repo):
    agile(repo)
    p = bl(repo, "--plan", "--write")
    assert p.returncode == 2 and "human-owned" in p.stdout
    p = bl(repo, "--plan", "--write", today="2026-10-14", owner=True)
    assert p.returncode == 2 and "may be planned from Fri 2026-10-16" in p.stdout
    p = bl(repo, "--plan", "--write", today="2026-10-17", owner=True)
    assert p.returncode == 2 and "is not one" in p.stdout
    p = bl(repo, "--plan", "--write", owner=True)
    assert p.returncode == 0, p.stdout
    ev = [json.loads(x) for x in (repo / "tracker" / "sprints.jsonl").read_text(encoding="utf-8").splitlines()]
    assert ev[0]["event"] == "planned" and ev[0]["sprint"] == "S1" and ev[0]["stories"] == ["US-01", "US-02"]
    assert ev[0]["start"] == "2026-10-16" and ev[0]["end"] == "2026-10-29" and ev[0]["capacity"] == 4
    again = bl(repo, "--plan", owner=True)
    assert again.returncode == 2 and "still open" in again.stdout


def test_a_story_walks_todo_doing_verify_rework_merge_review(repo):
    agile(repo)
    plan_sprint(repo)
    assert story_states(repo)["US-01"]["state"] == "TODO"
    p = bl(repo, "--start", "US-01")
    assert p.returncode == 0, p.stdout + p.stderr
    assert git(repo, "branch", "--show-current").strip() == "task/US-01-alpha-one"
    prompt = (repo / "tracker" / ".cache" / "prompt-US-01-P1.md").read_text(encoding="utf-8")
    assert "task/US-01-alpha-one" in prompt and "Alpha works" in prompt and "{{" not in prompt
    assert story_states(repo)["US-01"]["state"] == "DOING"
    put(repo, "tracker/handoff/US-01.md", "# US-01 handoff - claude, 2026-10-19\n## Plan\nthe plan\n")
    s = story_states(repo)["US-01"]
    assert s["state"] == "DOING" and any("Files changed" in n for n in s["notes"])
    put(repo, "tracker/handoff/US-01.md", AGILE_HANDOFF)
    assert story_states(repo)["US-01"]["state"] == "VERIFY"
    put(repo, "tracker/handoff/US-01-verify.md", "VERDICT: ACCEPT\nVerifier: Claude Opus\n")
    s = story_states(repo)["US-01"]
    assert s["state"] == "VERIFY" and any("same model family" in n for n in s["notes"])
    put(repo, "tracker/handoff/US-01-verify.md", "VERDICT: REJECT\nVerifier: Gemini\n1. fix x\n")
    assert story_states(repo)["US-01"]["state"] == "REWORK"
    put(repo, "tracker/handoff/US-01-verify.md", AGILE_VERIFY_OK)
    assert story_states(repo)["US-01"]["state"] == "MERGE"
    commit_all(repo, "feat(US-01): alpha")
    git(repo, "switch", "-q", "main")
    git(repo, "merge", "--ff-only", "task/US-01-alpha-one")
    git(repo, "branch", "-d", "task/US-01-alpha-one")
    assert story_states(repo)["US-01"]["state"] == "REVIEW"


def test_merge_steps_commit_a_verdict_that_is_not_committed_yet_and_only_then(repo):
    agile(repo)
    plan_sprint(repo)
    assert bl(repo, "--start", "US-01").returncode == 0
    put(repo, "tracker/handoff/US-01.md", AGILE_HANDOFF)
    commit_all(repo, "feat(US-01): alpha")
    put(repo, "tracker/handoff/US-01-verify.md", AGILE_VERIFY_OK)  # the verifier leaves it uncommitted
    out = bl(repo, today="2026-10-19").stdout
    assert "MERGE US-01" in out and 'git commit -m "chore(tracker): verdict for US-01"' in out, out
    assert out.index("verdict for US-01") < out.index("git switch main") < out.index("git merge --ff-only")
    commit_all(repo, "verdict")
    out = bl(repo, today="2026-10-19").stdout
    assert "MERGE US-01" in out and "verdict for US-01" not in out and "git switch main" in out, out


def test_a_verdict_swept_into_an_author_commit_does_not_block_the_merge_steps(repo):
    agile(repo)
    plan_sprint(repo)
    assert bl(repo, "--start", "US-01").returncode == 0
    put(repo, "tracker/handoff/US-01.md", AGILE_HANDOFF)
    put(repo, "tracker/handoff/US-01-verify.md", "VERDICT: REJECT\nVerifier: Gemini\n1. fix x\n")
    commit_all(repo, "feat(US-01): alpha (git add -A took the verdict too)")
    put(repo, "tracker/handoff/US-01-verify.md", AGILE_VERIFY_OK)  # the second verification overwrites a TRACKED file
    assert story_states(repo)["US-01"]["state"] == "MERGE"
    out = bl(repo, today="2026-10-19").stdout
    assert 'git commit -m "chore(tracker): verdict for US-01"' in out, out
    sw = subprocess.run(["git", "switch", "main"], cwd=repo, capture_output=True, text=True, encoding="utf-8")
    assert sw.returncode != 0 and "would be overwritten" in sw.stderr, sw.stderr  # the stall the printed first two steps prevent
    assert git(repo, "branch", "--show-current").strip() == "task/US-01-alpha-one"
    git(repo, "add", "tracker/handoff/US-01-verify.md")
    git(repo, "commit", "-q", "-m", "chore(tracker): verdict for US-01")
    git(repo, "switch", "-q", "main")
    git(repo, "merge", "--ff-only", "task/US-01-alpha-one")
    assert story_states(repo)["US-01"]["state"] == "REVIEW"


def test_start_refuses_a_second_story_a_dirty_tree_the_wrong_branch_and_unplanned_stories(repo):
    agile(repo)
    p = bl(repo, "--start", "US-01")
    assert p.returncode == 2 and "not in the active sprint" in p.stdout
    plan_sprint(repo)
    p = bl(repo, "--start", "US-04")
    assert p.returncode == 2 and "not in the active sprint" in p.stdout
    put(repo, "sidecar/dirty.py", "x = 1\n")
    p = bl(repo, "--start", "US-01")
    assert p.returncode == 2 and "not clean" in p.stdout
    (repo / "sidecar" / "dirty.py").unlink()
    git(repo, "switch", "-q", "-c", "side")
    p = bl(repo, "--start", "US-01")
    assert p.returncode == 2 and "start from main" in p.stdout
    git(repo, "switch", "-q", "main")
    assert bl(repo, "--start", "US-01").returncode == 0
    git(repo, "switch", "-q", "main")
    p = bl(repo, "--start", "US-02")
    assert p.returncode == 2 and "WIP limit 1" in p.stdout


def test_accept_needs_the_owner_review_state_level_and_a_real_note(repo):
    agile(repo, level=2)
    plan_sprint(repo)
    note = "watched it work on my phone"
    p = bl(repo, "--accept", "US-01", "--note", note)
    assert p.returncode == 2 and "human-owned" in p.stdout
    p = bl(repo, "--accept", "US-01", "--note", note, owner=True)
    assert p.returncode == 2 and "US-01 is TODO" in p.stdout
    merged_story(repo)
    p = bl(repo, "--accept", "US-01", "--note", note, owner=True)
    assert p.returncode == 2 and "below this story's target L3" in p.stdout
    put(repo, "tracker/status.json", json.dumps({"features": {"F-0001": {"level": 3}}}))
    commit_all(repo, "levels")
    p = bl(repo, "--accept", "US-01", "--note", "ok", owner=True)
    assert p.returncode == 2 and "at least 8 characters" in p.stdout
    p = bl(repo, "--accept", "US-01", "--note", note, owner=True)
    assert p.returncode == 0, p.stdout
    assert story_states(repo)["US-01"]["state"] == "DONE"
    assert bl(repo, "--accept", "US-01", "--note", note, owner=True).returncode == 2
    assert bl(repo, "--check").returncode == 0


def test_closing_a_sprint_needs_the_review_day_and_a_finished_retro_and_records_velocity(repo):
    agile(repo)
    plan_sprint(repo)
    merged_story(repo)
    assert bl(repo, "--accept", "US-01", "--note", "watched it work on my phone", owner=True).returncode == 0
    p = bl(repo, "--close-sprint", today="2026-10-20", owner=True)
    assert p.returncode == 2 and "review slot of sprint S1 is Tue 2026-10-27" in p.stdout
    p = bl(repo, "--close-sprint", today="2026-10-27", owner=True)
    assert p.returncode == 2 and "retro" in p.stdout.lower()
    p = bl(repo, "--retro", today="2026-10-27", owner=True)
    assert p.returncode == 0 and (repo / "tracker" / "retro" / "S1.md").is_file()
    p = bl(repo, "--close-sprint", today="2026-10-27", owner=True)
    assert p.returncode == 2 and "<fill>" in p.stdout
    put(repo, "tracker/retro/S1.md", "# Retro S1\n## Worked\n- the board\n## Did not work\n- long prompts\n## One change\none\ntwo\n")
    p = bl(repo, "--close-sprint", today="2026-10-27", owner=True)
    assert p.returncode == 2 and "exactly ONE line" in p.stdout
    put(repo, "tracker/retro/S1.md", "# Retro S1\n## Worked\n- the board\n## Did not work\n- long prompts\n## One change\nPaste the prompt file, not my own words\n")
    p = bl(repo, "--close-sprint", today="2026-10-27", owner=True)
    assert p.returncode == 0, p.stdout
    ev = [json.loads(x) for x in (repo / "tracker" / "sprints.jsonl").read_text(encoding="utf-8").splitlines()]
    closed = [e for e in ev if e["event"] == "closed"][0]
    assert closed["sprint"] == "S1" and closed["velocity"] == 2 and closed["done"] == ["US-01"] and closed["not_done"] == {"US-02": "TODO"}
    assert "Capacity for the next sprint: 4" in p.stdout


def test_capacity_follows_the_average_of_the_last_three_sprints(repo):
    from datetime import date, timedelta
    agile(repo)
    lines = []
    for i, v in enumerate((2, 3, 3), start=1):
        start = date(2026, 10, 16) + timedelta(days=14 * (i - 1))
        end = start + timedelta(days=13)
        lines.append(json.dumps({"event": "planned", "sprint": f"S{i}", "start": start.isoformat(), "end": end.isoformat(), "goal": "g",
                                 "capacity": 4, "points": 4, "stories": [], "date": start.isoformat()}))
        lines.append(json.dumps({"event": "closed", "sprint": f"S{i}", "date": end.isoformat(), "velocity": v, "done": [], "not_done": {}}))
    put(repo, "tracker/sprints.jsonl", "\n".join(lines) + "\n")
    commit_all(repo, "log")
    p = bl(repo, "--plan", today="2026-11-27")
    assert "capacity 2 points" in p.stdout and "average velocity" in p.stdout, p.stdout


def test_calendar_puts_planning_refinement_and_review_on_the_right_slots(repo):
    agile(repo)
    p = bl(repo, "--calendar")
    assert p.returncode == 0
    assert "Fri 2026-10-16 19:30  PLANNING" in p.stdout
    assert "Mon 2026-10-26 19:30  REFINEMENT" in p.stdout
    assert "Tue 2026-10-27 19:30  REVIEW+RETRO" in p.stdout
    assert "S2  Fri 2026-10-30 .. Thu 2026-11-12" in p.stdout


def test_refine_lists_the_decision_rows_and_writes_the_prompts(repo):
    agile(repo)
    p = bl(repo, "--refine")
    assert p.returncode == 0, p.stdout
    assert "D-10" in p.stdout and "US-05" in p.stdout
    p3 = (repo / "tracker" / ".cache" / "prompt-P3.md").read_text(encoding="utf-8")
    p5 = (repo / "tracker" / ".cache" / "prompt-P5.md").read_text(encoding="utf-8")
    assert "D-10" in p3 and "{{" not in p3
    assert "US-05" in p5 and "accept" in p5 and "{{" not in p5


def test_prompts_are_rendered_with_the_story_and_leave_no_placeholder(repo):
    agile(repo)
    plan_sprint(repo)
    p = bl(repo, "--prompt", "P1", "US-01")
    assert p.returncode == 0, p.stdout
    t1 = (repo / "tracker" / ".cache" / "prompt-US-01-P1.md").read_text(encoding="utf-8")
    assert "Alpha story" in t1 and "Alpha works" in t1 and "Run alpha" in t1 and "F-0001" in t1 and "{{" not in t1
    assert "Get-Content -Raw -Encoding UTF8" in p.stdout
    assert bl(repo, "--prompt", "P2", "US-01").returncode == 2
    assert bl(repo, "--start", "US-01").returncode == 0
    assert bl(repo, "--prompt", "P2", "US-01").returncode == 0
    t2 = (repo / "tracker" / ".cache" / "prompt-US-01-P2.md").read_text(encoding="utf-8")
    assert "VERDICT: ACCEPT" in t2 and "{{" not in t2
    assert bl(repo, "--prompt", "P9").returncode == 2
    p0 = bl(repo, "--prompt", "P0")
    assert p0.returncode == 0, p0.stdout
    t0 = (repo / "tracker" / ".cache" / "prompt-P0.md").read_text(encoding="utf-8")
    assert "You are an auditor" in t0 and "{{" not in t0 and not t0.startswith("# P0")


def test_the_p4_addendum_is_appended_when_a_story_asks_for_it(repo):
    agile(repo)
    put(repo, "tracker/backlog.yaml", AGILE_BACKLOG.replace("slug: alpha-one,", "slug: alpha-one, addendum: [P4],"))
    commit_all(repo, "addendum")
    assert bl(repo, "--prompt", "P1", "US-01").returncode == 0
    assert "Extra rules for this story" in (repo / "tracker" / ".cache" / "prompt-US-01-P1.md").read_text(encoding="utf-8")


def test_next_action_names_the_one_thing_to_do_each_day(repo):
    agile(repo)

    def nxt(day: str) -> list[str]:
        return json.loads(bl(repo, "--json", today=day).stdout)["next"]

    assert "planning day: Fri 2026-10-16" in " ".join(nxt("2026-10-12"))
    assert any(x.startswith("PLANNING") for x in nxt("2026-10-16"))
    plan_sprint(repo)
    assert any(x.startswith("START US-01") for x in nxt("2026-10-19"))
    assert "Not a slot day. Next slot: Fri 2026-10-23" in " ".join(nxt("2026-10-21"))
    assert any("REFINEMENT" in x for x in nxt("2026-10-26"))
    assert any(x.startswith("REVIEW + RETRO") for x in nxt("2026-10-27"))


def test_files_on_main_without_a_verdict_are_an_anomaly(repo):
    agile(repo)
    put(repo, "tracker/handoff/US-02.md", AGILE_HANDOFF.replace("US-01", "US-02"))
    commit_all(repo, "oops")
    assert story_states(repo)["US-02"]["state"] == "ANOMALY"
    assert any("US-02" in x for x in bl(repo, "--check").stdout.splitlines() if x.startswith("WARN"))


def test_check_catches_an_accepted_event_for_a_story_nobody_verified(repo):
    agile(repo)
    put(repo, "tracker/sprints.jsonl", json.dumps({"event": "accepted", "story": "US-01", "date": "2026-10-20", "note": "typed by hand"}) + "\n")
    p = bl(repo, "--check")
    assert p.returncode == 1 and "US-01 is recorded as accepted but is TODO" in p.stdout


def test_check_catches_two_stories_in_progress(repo):
    agile(repo)
    git(repo, "branch", "task/US-01-alpha-one")
    git(repo, "branch", "task/US-02-gamma-two")
    p = bl(repo, "--check")
    assert p.returncode == 1 and "WIP limit 1 exceeded" in p.stdout


def test_the_hooks_reminder_appears_only_when_the_guard_was_clean(repo):
    agile(repo)
    put(repo, "tracker/.cache/guard_report.json", json.dumps({"verdict": "FAIL", "mode": "full", "only": None, "counts": {"FAIL": 2, "WARN": 0, "INFO": 0}, "findings": []}))
    assert not any("HOOKS ON" in x for x in json.loads(bl(repo, "--json").stdout)["next"])
    put(repo, "tracker/.cache/guard_report.json", json.dumps({"verdict": "PASS", "mode": "fast", "only": None, "counts": {"FAIL": 0, "WARN": 0, "INFO": 0}, "findings": []}))
    assert not any("HOOKS ON" in x for x in json.loads(bl(repo, "--json").stdout)["next"]), "a fast run must not trigger the reminder"
    put(repo, "tracker/.cache/guard_report.json", json.dumps({"verdict": "PASS", "mode": "full", "only": ["secrets"], "counts": {"FAIL": 0, "WARN": 0, "INFO": 0}, "findings": []}))
    assert not any("HOOKS ON" in x for x in json.loads(bl(repo, "--json").stdout)["next"]), "an --only run must not trigger the reminder"
    put(repo, "tracker/.cache/guard_report.json", json.dumps({"verdict": "PASS", "mode": "full", "only": None, "counts": {"FAIL": 0, "WARN": 0, "INFO": 0}, "findings": []}))
    assert any("HOOKS ON" in x for x in json.loads(bl(repo, "--json").stdout)["next"])
    git(repo, "config", "core.hooksPath", ".githooks")
    assert not any("HOOKS ON" in x for x in json.loads(bl(repo, "--json").stdout)["next"])


def test_the_bundles_own_backlog_is_valid():
    if not (BUNDLE / "tracker" / "backlog.yaml").is_file():
        pytest.skip("backlog not installed yet")
    p = subprocess.run([PY, str(TOOLS / "tracker" / "backlog.py"), "--root", str(BUNDLE), "--check"], capture_output=True, text=True,
                       encoding="utf-8", env=base_env())
    assert p.returncode == 0, p.stdout


def test_only_rejects_unknown_check_names_and_promotes_full_only_checks(repo):
    p = subprocess.run([PY, str(repo / "tools" / "guard" / "check.py"), "--only", "secrets,gitignroe"], cwd=repo, capture_output=True,
                       text=True, encoding="utf-8", env=base_env())
    assert p.returncode == 2 and "unknown check name" in p.stdout and "gitignroe" in p.stdout
    r = guard(repo, "--only", "gitignore")
    assert r.report["mode"] == "full" and r.report["only"] == ["gitignore"], r.report


def test_a_partial_guard_report_is_not_evidence_for_the_tracker(repo):
    put(repo, "tracker/.cache/guard_report.json", json.dumps({"verdict": "PASS", "mode": "fast", "only": None,
                                                              "counts": {"FAIL": 0, "WARN": 0, "INFO": 0}, "findings": []}))
    p = tracker(repo, "verify.py", "--static", "--no-guard")
    assert p.returncode == 0, p.stdout + p.stderr
    assert json.loads((repo / "tracker" / "status.json").read_text(encoding="utf-8"))["guard"] is None


# ================================================================== 17. the autopilot (tools/tracker/autopilot.py), driven by a fake `agy`
FAKE_AGY = r'''
import json, os, re, subprocess, sys
from pathlib import Path

args = sys.argv[1:]
if args[:1] == ["models"]:
    print("claude-test-1\ngemini-test-1")
    sys.exit(0)
prompt = args[args.index("-p") + 1]
model = args[args.index("--model") + 1]
mode = os.environ.get("FAKE_AGY_MODE", "good")
if "Do not run any command and do not create or edit any file." in prompt:  # the driver's smoke question
    with open(os.environ["FAKE_AGY_LOG"] + ".smoke", "a", encoding="utf-8") as f:
        f.write(json.dumps({"model": model, "key": os.environ.get("MY_API_KEY"), "timeout_arg": args[args.index("--print-timeout") + 1]}) + "\n")
    if mode == "smoke-hang":
        import time
        time.sleep(60)
    if mode == "fail":
        print("boom", file=sys.stderr)
        sys.exit(3)
    answer = re.search(r"answer exactly `([^`]+)`", (Path.cwd() / "AGENTS.md").read_text(encoding="utf-8")).group(1)
    if mode == "noisy":
        print("agy: a new version is available")
    print(json.dumps({"conversation_id": "smoke", "status": "SUCCESS", "response": "no idea" if mode == "no-canary" else answer}))
    sys.exit(0)
state = Path(os.environ["FAKE_AGY_STATE"])
n = (int(state.read_text()) if state.exists() else 0) + 1
state.write_text(str(n))
verifier = "You verify branch" in prompt
with open(os.environ["FAKE_AGY_LOG"], "a", encoding="utf-8") as f:
    f.write(json.dumps({"n": n, "model": model, "kind": "verify" if verifier else "author", "owner_var": os.environ.get("LOCUS_ALLOW_GUARD_EDIT"),
                        "key": os.environ.get("MY_API_KEY"), "autopilot_rules": "AUTOPILOT MODE" in prompt}) + "\n")
root = Path.cwd()
canary = re.search(r"answer exactly `([^`]+)`", (root / "AGENTS.md").read_text(encoding="utf-8")).group(1)
sid = re.search(r"story (US-\d+)", prompt).group(1)


def reply(text):
    head = "" if mode == "no-canary" else "CANARY: " + canary + "\n"
    if mode == "noisy":
        print("agy: a new version is available")
    print(json.dumps({"conversation_id": "c%d" % n, "status": "SUCCESS", "response": head + text, "duration_seconds": 1, "num_turns": 1}))
    sys.exit(0)


if mode == "fail":
    print("boom", file=sys.stderr)
    sys.exit(3)
handoff_dir = root / "tracker" / "handoff"
handoff_dir.mkdir(parents=True, exist_ok=True)
if verifier:
    vf = state.with_suffix(".v")
    v = (int(vf.read_text()) if vf.exists() else 0) + 1
    vf.write_text(str(v))
    verdict = "REJECT" if (mode == "reject-always" or (mode == "reject-once" and v == 1)) else "ACCEPT"
    (handoff_dir / ("%s-verify.md" % sid)).write_text("VERDICT: %s\nVerifier: %s\n## Findings\n%s" % (verdict, model, "1. fix the thing\n" if verdict == "REJECT" else "none\n"), encoding="utf-8")
    reply("VERIFIED %s: %s" % (sid, verdict))
if mode == "noop":
    reply("DONE " + sid)
q = handoff_dir / ("%s-questions.md" % sid)
if mode == "questions" and not q.exists():
    q.write_text("Q1: which port?\nRecommended: 8420\nAnswer:\n", encoding="utf-8")
    reply("QUESTIONS " + sid)
heads = ["Plan (written before coding)", "Files changed (every path, one line each)", "Spec sentences implemented", "Config keys and defaults added",
         "Commands run, with real output", "Open questions / errata rows touched", "Discard table", "Security", "Not done"]
(handoff_dir / ("%s.md" % sid)).write_text("# %s handoff - %s, 2026-10-19\n" % (sid, model) + "".join("## %s\ncall %d: ok\n" % (h, n) for h in heads), encoding="utf-8")
if mode == "bad-layout":
    (root / "core").mkdir(exist_ok=True)
    (root / "core" / "x.py").write_text("x = 1\n", encoding="utf-8")
subprocess.run(["git", "add", "-A"], check=True, capture_output=True)
subprocess.run(["git", "commit", "-q", "-m", "feat(%s): work %d" % (sid, n)], check=True, capture_output=True)
reply("DONE " + sid)
'''


def ap_ready(repo: Path, plan: bool = True, backlog: str | None = None) -> None:
    agile(repo)
    put(repo, "AGENTS.md", "<!-- LOCUS-AGENTS-V1 -->\n# Locus rules\nIf the human asks for the canary, answer exactly `blue-heron-47`.\n")
    if backlog:
        put(repo, "tracker/backlog.yaml", backlog)
    commit_all(repo, "agents file")
    if plan:
        plan_sprint(repo)
    git(repo, "config", "core.hooksPath", ".githooks")


def autopilot(repo: Path, tmp_path: Path, *extra: str, mode: str = "good", env: dict | None = None, models: tuple = ("claude-test-1", "gemini-test-1"),
              today: str = "2026-10-19") -> subprocess.CompletedProcess:
    script = tmp_path / "fake_agy.py"
    script.write_text(FAKE_AGY, encoding="utf-8")
    e = {"FAKE_AGY_STATE": str(tmp_path / "calls.n"), "FAKE_AGY_LOG": str(tmp_path / "calls.jsonl"), "FAKE_AGY_MODE": mode, **(env or {})}
    return tracker(repo, "autopilot.py", "--agy", f"{PY} {script}", "--author-model", models[0], "--verifier-model", models[1], "--today", today, *extra, env=e)


def agy_calls(tmp_path: Path) -> list[dict]:
    p = tmp_path / "calls.jsonl"
    return [json.loads(x) for x in p.read_text(encoding="utf-8").splitlines()] if p.exists() else []


def smoke_calls(tmp_path: Path) -> list[dict]:
    p = tmp_path / "calls.jsonl.smoke"
    return [json.loads(x) for x in p.read_text(encoding="utf-8").splitlines()] if p.exists() else []


def test_autopilot_walks_every_planned_story_to_review_and_stops_at_the_owner_gate(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--max-steps", "12")
    assert p.returncode == 0, p.stdout + p.stderr
    assert "nothing is left to start in sprint S1" in p.stdout and "US-01, US-02" in p.stdout
    st = story_states(repo)
    assert st["US-01"]["state"] == "REVIEW" and st["US-02"]["state"] == "REVIEW"
    assert git(repo, "branch", "--show-current").strip() == "main"
    assert git(repo, "branch", "--list", "task/*").strip() == ""
    cs = agy_calls(tmp_path)
    assert [(c["kind"], c["model"]) for c in cs] == [("author", "claude-test-1"), ("verify", "gemini-test-1")] * 2
    assert all(c["autopilot_rules"] and c["owner_var"] is None for c in cs)
    log = git(repo, "log", "--oneline")
    assert "refresh after US-01" in log and "refresh after US-02" in log
    assert (repo / "tracker" / ".cache" / "autopilot.log").is_file()
    assert bl(repo, "--accept", "US-01", "--note", "I watched it", today="2026-10-27").returncode == 2  # accepting stays the owner's


def test_autopilot_stops_on_questions_and_carries_on_once_the_owner_has_answered(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, mode="questions")
    assert p.returncode == 0 and "questions only you can answer" in p.stdout, p.stdout + p.stderr
    q = repo / "tracker" / "handoff" / "US-01-questions.md"
    assert q.is_file() and len(agy_calls(tmp_path)) == 1 and story_states(repo)["US-01"]["state"] == "DOING"
    p = autopilot(repo, tmp_path, mode="questions")
    assert p.returncode == 0 and "waiting for your answers" in p.stdout and len(agy_calls(tmp_path)) == 1
    q.write_text(q.read_text(encoding="utf-8").replace("Answer:", "Answer: 8420"), encoding="utf-8")
    p = autopilot(repo, tmp_path, "--max-steps", "3")
    assert p.returncode == 0, p.stdout + p.stderr
    assert story_states(repo)["US-01"]["state"] == "REVIEW"


def test_autopilot_sends_a_rejected_story_back_and_gives_up_after_two_rounds(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--max-steps", "5", mode="reject-once")
    assert p.returncode == 0 and "step budget is used" in p.stdout, p.stdout + p.stderr
    assert story_states(repo)["US-01"]["state"] == "REVIEW"
    assert [c["kind"] for c in agy_calls(tmp_path)] == ["author", "verify", "author", "verify"]


def test_autopilot_does_not_loop_on_a_story_the_verifier_keeps_rejecting(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--max-steps", "30", mode="reject-always")
    assert p.returncode == 1 and "after 2 rounds of rework" in p.stdout, p.stdout + p.stderr
    assert len(agy_calls(tmp_path)) == 6 and story_states(repo)["US-01"]["state"] == "REWORK"


def test_autopilot_stops_when_the_agent_skips_the_canary_or_fails(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--smoke-seconds", "0", mode="no-canary")
    assert p.returncode == 1 and "did not load AGENTS.md" in p.stdout, p.stdout + p.stderr
    p = autopilot(repo, tmp_path, "--smoke-seconds", "0", mode="fail")
    assert p.returncode == 1 and "agy exited 3" in p.stdout, p.stdout + p.stderr


def test_autopilot_smoke_test_asks_both_models_for_the_canary_before_the_first_step(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--max-steps", "1")
    assert p.returncode == 0, p.stdout + p.stderr
    sm = smoke_calls(tmp_path)
    assert [c["model"] for c in sm] == ["claude-test-1", "gemini-test-1"] and all(c["timeout_arg"] == "3m" for c in sm)
    assert len(agy_calls(tmp_path)) == 1


def test_autopilot_smoke_test_stops_a_hanging_or_ruleless_agy_before_any_work(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--smoke-seconds", "2", mode="smoke-hang")
    assert p.returncode == 2 and "smoke test failed" in p.stdout and "did not answer within 2 seconds" in p.stdout, p.stdout + p.stderr
    assert "manual flow" in p.stdout
    p = autopilot(repo, tmp_path, mode="no-canary")
    assert p.returncode == 2 and "smoke test failed" in p.stdout and "did not load AGENTS.md" in p.stdout, p.stdout + p.stderr
    p = autopilot(repo, tmp_path, mode="fail")
    assert p.returncode == 2 and "smoke test failed" in p.stdout and "agy exited 3" in p.stdout, p.stdout + p.stderr
    assert agy_calls(tmp_path) == [] and git(repo, "branch", "--list", "task/*").strip() == ""
    assert git(repo, "branch", "--show-current").strip() == "main"


def test_autopilot_reads_the_json_envelope_even_when_agy_prints_a_banner_around_it(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--max-steps", "1", mode="noisy")
    assert p.returncode == 0, p.stdout + p.stderr
    assert story_states(repo)["US-01"]["state"] == "VERIFY" and len(agy_calls(tmp_path)) == 1


def test_autopilot_stops_when_an_agent_changes_nothing_twice(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--max-steps", "10", mode="noop")
    assert p.returncode == 1 and "no progress after 2 attempts" in p.stdout, p.stdout + p.stderr
    assert len(agy_calls(tmp_path)) == 2


def test_autopilot_refuses_to_start_without_hooks_with_one_model_family_or_with_the_owner_variable(repo, tmp_path):
    ap_ready(repo)
    git(repo, "config", "--unset", "core.hooksPath")
    p = autopilot(repo, tmp_path)
    assert p.returncode == 2 and "commit hooks are off" in p.stdout
    git(repo, "config", "core.hooksPath", ".githooks")
    p = autopilot(repo, tmp_path, models=("claude-a", "claude-b"))
    assert p.returncode == 2 and "different model families" in p.stdout
    p = autopilot(repo, tmp_path, env={"LOCUS_ALLOW_GUARD_EDIT": "1"})
    assert p.returncode == 2 and "owner variable" in p.stdout
    p = autopilot(repo, tmp_path, models=("claude-test-1", "gemini-missing"))
    assert p.returncode == 2 and "not in the output of" in p.stdout
    p = autopilot(repo, tmp_path, "--agy", "no-such-agy-program")
    assert p.returncode == 2 and "was not found" in p.stdout
    assert agy_calls(tmp_path) == []


def test_autopilot_refuses_to_start_on_a_red_guard(repo, tmp_path):
    ap_ready(repo)
    put(repo, "core/x.py", "x = 1\n")
    p = autopilot(repo, tmp_path)
    assert p.returncode == 2 and "guard shows FAIL" in p.stdout
    assert agy_calls(tmp_path) == []


def test_autopilot_dry_run_the_stop_file_and_the_sprint_gates_call_no_agent(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--dry-run")
    assert p.returncode == 0 and "DRY RUN" in p.stdout and "START US-01" in p.stdout, p.stdout + p.stderr
    assert agy_calls(tmp_path) == [] and git(repo, "branch", "--show-current").strip() == "main"
    assert git(repo, "branch", "--list", "task/*").strip() == ""
    put(repo, "tracker/.cache/STOP", "")
    p = autopilot(repo, tmp_path)
    assert p.returncode == 0 and "STOP file exists" in p.stdout
    (repo / "tracker" / ".cache" / "STOP").unlink()
    p = autopilot(repo, tmp_path, today="2026-10-30")
    assert p.returncode == 0 and "sprint S1 ended" in p.stdout
    assert agy_calls(tmp_path) == []


def test_autopilot_leaves_a_hands_on_story_to_the_owner_and_carries_on_after_it(repo, tmp_path):
    hands_on = AGILE_BACKLOG.replace('title: "Alpha story", points: 2,', 'title: "Alpha story", autopilot: false, points: 2,')
    ap_ready(repo, backlog=hands_on)
    p = autopilot(repo, tmp_path)
    assert p.returncode == 0 and "US-01 is marked 'autopilot: false'" in p.stdout, p.stdout + p.stderr
    assert agy_calls(tmp_path) == [] and git(repo, "branch", "--list", "task/*").strip() == ""
    merged_story(repo)  # the owner does US-01 by hand, with the board
    p = autopilot(repo, tmp_path, "--max-steps", "2")
    assert p.returncode == 0 and "step budget is used" in p.stdout, p.stdout + p.stderr
    st = story_states(repo)
    assert st["US-01"]["state"] == "REVIEW" and st["US-02"]["state"] == "MERGE"
    put(repo, "tracker/backlog.yaml", AGILE_BACKLOG.replace('title: "Alpha story", points: 2,', 'title: "Alpha story", autopilot: "no", points: 2,'))
    p = bl(repo, "--check")
    assert p.returncode == 1 and "US-01: autopilot must be true or false" in p.stdout, p.stdout


def test_autopilot_does_not_start_before_the_owner_has_planned_a_sprint(repo, tmp_path):
    ap_ready(repo, plan=False)
    p = autopilot(repo, tmp_path)
    assert p.returncode == 0 and "planning is yours" in p.stdout
    assert agy_calls(tmp_path) == []


def test_autopilot_keeps_secret_looking_variables_away_from_the_agent(repo, tmp_path):
    ap_ready(repo)
    autopilot(repo, tmp_path, "--max-steps", "1", env={"MY_API_KEY": "should-not-leak"})
    cs = agy_calls(tmp_path)
    assert len(cs) == 1 and cs[0]["key"] is None and all(c["key"] is None for c in smoke_calls(tmp_path))
    autopilot(repo, tmp_path, "--max-steps", "1", "--pass-env", "MY_API_KEY", env={"MY_API_KEY": "passed-on-purpose"})
    assert agy_calls(tmp_path)[-1]["key"] == "passed-on-purpose"


def test_autopilot_does_not_merge_when_the_guard_fails(repo, tmp_path):
    ap_ready(repo)
    p = autopilot(repo, tmp_path, "--max-steps", "6", mode="bad-layout")
    assert p.returncode == 1 and "the guard shows FAIL" in p.stdout, p.stdout + p.stderr
    assert story_states(repo)["US-01"]["state"] == "MERGE"
    assert git(repo, "branch", "--show-current").strip() == "task/US-01-alpha-one"
    assert "core" not in git(repo, "ls-tree", "--name-only", "main").split()
